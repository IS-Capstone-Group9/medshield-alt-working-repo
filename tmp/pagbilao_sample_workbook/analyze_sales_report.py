from __future__ import annotations

import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path(r"C:\Users\Adrian\projects\medshield1\medshield\data\medshield\dataset_xlsx\Sales Report.xlsx")
OUT = Path(r"C:\Users\Adrian\projects\medshield1\medshield\tmp\pagbilao_sample_workbook\sales_report_analysis.json")

wb = load_workbook(SOURCE, read_only=True, data_only=True)
result: dict[str, object] = {"sheet_names": wb.sheetnames, "sheets": []}

contract_patterns = [
    r"\bPAGBILAO\b",
    r"\bPESO\b",
    r"\bQMC\b",
    r"\bCONTRACT\b",
    r"\bPURCHASE\s*(?:ORDER|REQUEST|REQ\.?|NO\.?|#)",
    r"\bPO\s*(?:NO\.?|#|NUMBER)",
    r"\bPR\s*(?:NO\.?|#|NUMBER)",
    r"\bDR\s*(?:NO\.?|#|NUMBER)",
    r"\bDELIVERY\s*(?:RECEIPT|ORDER)",
    r"\bBIDDING\b",
    r"\bPROJECT\b",
]
contract_re = re.compile("|".join(contract_patterns), re.IGNORECASE)

for ws in wb.worksheets:
    sample = []
    header_row = None
    header = None
    for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = list(row)
        if len(sample) < 8:
            sample.append(values[:15])
        normalized = [str(value).strip().lower() if value is not None else "" for value in values]
        if header_row is None and "area" in normalized and any(value in normalized for value in ("product", "product name", "description")):
            header_row = row_index
            header = values
    sheet_info = {
        "title": ws.title,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "header_row": header_row,
        "header": header,
        "sample": sample,
    }
    result["sheets"].append(sheet_info)

# Detailed analysis on sheets that contain transaction columns.
transactions = []
for ws in wb.worksheets:
    rows = ws.iter_rows(values_only=True)
    header = None
    header_row = None
    for row_index, row in enumerate(rows, start=1):
        normalized = [str(value).strip().lower() if value is not None else "" for value in row]
        if "area" in normalized and any(value in normalized for value in ("product", "product name", "description")):
            header = list(row)
            header_row = row_index
            break
    if not header:
        continue
    keys = [str(value).strip().lower().replace(" ", "_") if value is not None else f"column_{i+1}" for i, value in enumerate(header)]
    for source_row, row in enumerate(rows, start=header_row + 1):
        values = list(row)
        record = {keys[i]: values[i] if i < len(values) else None for i in range(len(keys))}
        area = str(record.get("area") or "").strip()
        product_key = next((key for key in ("product", "product_name", "description") if key in record), None)
        product = str(record.get(product_key) or "").strip() if product_key else ""
        if not area and not product:
            continue
        record["_sheet"] = ws.title
        record["_source_row"] = source_row
        record["_product"] = product
        record["_contract_like"] = bool(contract_re.search(product)) or product == "" or product.upper() in {"#REF!", "N/A", "NA", "NONE", "-"}
        transactions.append(record)

product_counts = Counter(item["_product"] for item in transactions)
contract_rows = [item for item in transactions if item["_contract_like"]]
by_sheet = {}
for sheet_name in wb.sheetnames:
    sheet_rows = [item for item in transactions if item["_sheet"] == sheet_name]
    sheet_contracts = [item for item in sheet_rows if item["_contract_like"]]
    prefix_counts = Counter()
    for item in sheet_contracts:
        product = item["_product"].upper()
        if not product:
            prefix = "BLANK"
        elif product.startswith("PAGBILAO"):
            prefix = "PAGBILAO"
        elif product.startswith("PESO"):
            prefix = "PESO"
        elif product.startswith("QMC"):
            prefix = "QMC"
        elif product == "#REF!":
            prefix = "#REF!"
        else:
            prefix = "OTHER_CONTRACT_LIKE"
        prefix_counts[prefix] += 1
    by_sheet[sheet_name] = {
        "rows": len(sheet_rows),
        "blank_product_rows": sum(not item["_product"] for item in sheet_rows),
        "contract_like_rows": len(sheet_contracts),
        "contract_prefix_counts": dict(prefix_counts),
        "contract_like_products": Counter(item["_product"] for item in sheet_contracts).most_common(),
    }
result["transactions"] = {
    "count": len(transactions),
    "product_unique": len(product_counts),
    "blank_product_rows": sum(not item["_product"] for item in transactions),
    "contract_like_rows": len(contract_rows),
    "contract_like_product_counts": Counter(item["_product"] for item in contract_rows).most_common(),
    "contract_like_examples": contract_rows[:100],
    "by_sheet": by_sheet,
}

OUT.write_text(json.dumps(result, ensure_ascii=False, default=str, indent=2), encoding="utf-8")
print(json.dumps({
    "sheet_names": wb.sheetnames,
    "sheet_count": len(wb.sheetnames),
    "transaction_count": len(transactions),
    "unique_products": len(product_counts),
    "blank_product_rows": result["transactions"]["blank_product_rows"],
    "contract_like_rows": len(contract_rows),
    "analysis_path": str(OUT),
}, indent=2))
