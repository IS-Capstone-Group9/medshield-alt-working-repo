from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

PATH = Path(r"C:\Users\Adrian\projects\medshield1\medshield\outputs\full_medshield_2025_allocation\Full_MedShield_2025_with_all_contract_medical_product_allocation.xlsx")
wb = load_workbook(PATH, read_only=True, data_only=True)
ws = wb["Full 2025 Dataset"]
headers = [cell.value for cell in next(ws.iter_rows(min_row=6, max_row=6))]
index = {name: position for position, name in enumerate(headers)}
rows = []
for values in ws.iter_rows(min_row=7, max_row=4576, values_only=True):
    rows.append({name: values[position] for name, position in index.items()})

contract_re = re.compile(r"^(?:PAGBILAO|QMC|PESO|ABOITIZ|GULANG GULANG|MARINDUQUE|PHO|PPDC|PPOC|PROVINCIAL TOURISM OFFICE|TOURISM)\s*#", re.I)
children = [row for row in rows if row["record_type"] == "estimated_medical_product_child"]
result = {
    "sheets": wb.sheetnames,
    "full_rows": len(rows),
    "unique_source_record_ids": len({row["source_record_id"] for row in rows}),
    "blank_products": sum(not str(row["product"] or "").strip() for row in rows),
    "contract_labels_in_product": sum(bool(contract_re.search(str(row["product"] or ""))) for row in rows),
    "estimated_children": len(children),
    "children_missing_provisional_sku": sum(not row["provisional_sku"] for row in children),
    "children_outside_medical_groups": sum(row["product_group"] not in {"medicine", "medical_device", "medical_supply"} for row in children),
    "totals": {
        "total_cp": round(sum(float(row["total_cp"] or 0) for row in rows), 2),
        "discount": round(sum(float(row["discount"] or 0) for row in rows), 2),
        "net_cp": round(sum(float(row["net_cp"] or 0) for row in rows), 2),
        "total_tp": round(sum(float(row["total_tp"] or 0) for row in rows), 2),
        "net_income": round(sum(float(row["net_income"] or 0) for row in rows), 2),
    },
}
assert result["full_rows"] == 4570, result
assert result["unique_source_record_ids"] == 4570, result
assert result["blank_products"] == 0, result
assert result["contract_labels_in_product"] == 0, result
assert result["estimated_children"] == 763, result
assert result["children_missing_provisional_sku"] == 0, result
assert result["children_outside_medical_groups"] == 0, result
assert result["totals"] == {
    "total_cp": 185777264.24,
    "discount": 380443.20,
    "net_cp": 185396821.04,
    "total_tp": 94760630.25,
    "net_income": 90636190.79,
}, result
print(json.dumps(result, indent=2))
