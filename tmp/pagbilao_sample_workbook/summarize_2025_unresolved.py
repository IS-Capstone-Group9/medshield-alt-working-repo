from __future__ import annotations

import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

SOURCE = Path(r"C:\Users\Adrian\projects\medshield1\medshield\data\medshield\dataset_xlsx\Sales Report.xlsx")
OUT = Path(r"C:\Users\Adrian\projects\medshield1\medshield\tmp\pagbilao_sample_workbook\unresolved_2025_summary.json")

wb = load_workbook(SOURCE, read_only=True, data_only=True)
ws = wb["2025"]
headers = [str(value).strip() if value is not None else "" for value in next(ws.iter_rows(min_row=6, max_row=6, values_only=True))]

def num(value):
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").replace("₱", "").replace("%", "").strip())
    except ValueError:
        return 0.0

rows = []
for source_row, values in enumerate(ws.iter_rows(min_row=7, values_only=True), start=7):
    record = dict(zip(headers, values))
    if not str(record.get("Area") or "").strip() and not str(record.get("DR Number") or "").strip():
        continue
    product = str(record.get("Product") or "").strip()
    upper = product.upper()
    if upper.startswith("PAGBILAO"):
        kind = "PAGBILAO"
    elif upper.startswith("QMC"):
        kind = "QMC"
    elif upper.startswith("PESO"):
        kind = "PESO"
    elif not product:
        kind = "BLANK"
    else:
        continue
    rows.append({"source_row": source_row, "kind": kind, **record})

summary = {}
for kind in ["PAGBILAO", "QMC", "PESO", "BLANK"]:
    items = [row for row in rows if row["kind"] == kind]
    summary[kind] = {
        "row_count": len(items),
        "areas": dict(Counter(str(row.get("Area") or "") for row in items)),
        "missing_dr": sum(not str(row.get("DR Number") or "").strip() for row in items),
        "missing_date": sum(row.get("Date Delivered") in (None, "") for row in items),
        "nonpositive_net_cp": sum(num(row.get("Net CP")) <= 0 for row in items),
        "nonpositive_total_tp": sum(num(row.get("TOTAL TP")) <= 0 for row in items),
        "total_cp": round(sum(num(row.get("Total CP")) for row in items), 2),
        "discount": round(sum(num(row.get("Disc")) for row in items), 2),
        "net_cp": round(sum(num(row.get("Net CP")) for row in items), 2),
        "total_tp": round(sum(num(row.get("TOTAL TP")) for row in items), 2),
        "net_income": round(sum(num(row.get("Net Income")) for row in items), 2),
        "sample": items[:20],
    }

summary["all"] = {
    "row_count": len(rows),
    "total_cp": round(sum(num(row.get("Total CP")) for row in rows), 2),
    "discount": round(sum(num(row.get("Disc")) for row in rows), 2),
    "net_cp": round(sum(num(row.get("Net CP")) for row in rows), 2),
    "total_tp": round(sum(num(row.get("TOTAL TP")) for row in rows), 2),
    "net_income": round(sum(num(row.get("Net Income")) for row in rows), 2),
}

OUT.write_text(json.dumps(summary, default=str, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(summary, default=str, ensure_ascii=False, indent=2))
