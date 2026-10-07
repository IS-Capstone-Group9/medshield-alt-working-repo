from __future__ import annotations

import csv
import hashlib
import json
import math
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(r"C:\Users\Adrian\projects\medshield1\medshield")
SOURCE = BASE / "data" / "medshield" / "dataset_xlsx" / "Sales Report.xlsx"
CLASSIFICATION = BASE / "outputs" / "product_grouping_2017_2025" / "product_classification_candidates.csv"
OUT = BASE / "tmp" / "pagbilao_sample_workbook" / "full_2025_all_contracts_staging.json"
MEDICAL_GROUPS = {"medicine", "medical_device", "medical_supply"}
INVALID_PRODUCTS = {"", "#REF!", "N/A", "NA", "NONE", "-"}


def number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("₱", "").replace("%", "")
    if text in {"", "-"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def iso_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if value is None:
        return ""
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%m-%d-%Y"):
        try:
            return datetime.strptime(text[:10], pattern).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return ""


def normalized_product(value):
    return re.sub(r"\s+", " ", str(value or "").strip().upper())


def provisional_sku(product):
    digest = hashlib.sha1(normalized_product(product).encode("utf-8")).hexdigest()[:12].upper()
    return f"MED-{digest}"


classification = {}
with CLASSIFICATION.open(encoding="utf-8-sig", newline="") as handle:
    for row in csv.DictReader(handle):
        classification[normalized_product(row["raw_product"])] = row


def class_info(product):
    key = normalized_product(product)
    row = classification.get(key)
    if not row:
        return {
            "group": "unresolved",
            "category": "unclassified_product_candidate",
            "confidence": "low",
            "rule": "no_governed_rule_match",
            "canonical_name": str(product or "").strip(),
        }
    return {
        "group": row["product_group"],
        "category": row["product_category"],
        "confidence": row["confidence"],
        "rule": row["classification_rule"],
        "canonical_name": row.get("proposed_canonical_sku") or row.get("normalized_product") or str(product or "").strip(),
    }


def read_sheet(ws, year):
    headers = [str(value).strip() if value is not None else "" for value in next(ws.iter_rows(min_row=6, max_row=6, values_only=True))]
    rows = []
    for source_row, values in enumerate(ws.iter_rows(min_row=7, values_only=True), start=7):
        raw = dict(zip(headers, values))
        area = str(raw.get("Area") or "").strip()
        dr_number = str(raw.get("DR Number") or "").strip()
        product = str(raw.get("Product") or "").strip()
        if not area and not dr_number and not product:
            continue
        rows.append({
            "year": year,
            "source_row": source_row,
            "area": area,
            "dr_number": dr_number,
            "date_delivered": iso_date(raw.get("Date Delivered")),
            "product": product,
            "qty": number(raw.get("Qty")),
            "cp": number(raw.get("CP")),
            "total_cp": number(raw.get("Total CP")),
            "discount": number(raw.get("Disc")) or 0.0,
            "net_cp": number(raw.get("Net CP")),
            "tp_unit": number(raw.get("TP/UNIT")),
            "total_tp": number(raw.get("TOTAL TP")),
            "net_income": number(raw.get("Net Income")),
            "margin_pct_source": number(raw.get("%")),
        })
    return rows


wb = load_workbook(SOURCE, read_only=True, data_only=True)
rows_by_year = {int(name): read_sheet(wb[name], int(name)) for name in wb.sheetnames if name.isdigit()}
source_2025 = rows_by_year[2025]


def replacement_type(row):
    product = normalized_product(row["product"])
    info = class_info(row["product"])
    if product.startswith("PAGBILAO"):
        return "pagbilao_contract"
    if product.startswith("QMC"):
        return "qmc_contract"
    if product.startswith("PESO"):
        return "peso_contract"
    if product in INVALID_PRODUCTS:
        return "blank_or_invalid_product"
    if info["group"] == "non_product":
        return "other_non_product"
    return ""


def normalized_source_row(row):
    info = class_info(row["product"])
    net = row["net_cp"] or 0.0
    cost = row["total_tp"] or 0.0
    gross = row["net_income"] if row["net_income"] is not None else net - cost
    margin = gross / net if net else None
    source_id = f"SRC-XLSX-2025-R{row['source_row']:04d}"
    return {
        "area": row["area"],
        "dr_number": row["dr_number"],
        "date_delivered": row["date_delivered"],
        "product": row["product"],
        "qty": row["qty"],
        "cp": row["cp"],
        "total_cp": row["total_cp"],
        "discount": row["discount"],
        "net_cp": row["net_cp"],
        "tp_unit": row["tp_unit"],
        "total_tp": row["total_tp"],
        "net_income": gross,
        "margin_pct": margin,
        "source_year": 2025,
        "source_sheet": "2025",
        "source_workbook_row": row["source_row"],
        "source_record_id": source_id,
        "parent_source_record_id": "",
        "child_sequence": None,
        "original_product_label": row["product"],
        "replacement_type": "",
        "record_type": "original_product_row",
        "product_group": info["group"],
        "product_category": info["category"],
        "classification_confidence": info["confidence"],
        "classification_rule": info["rule"],
        "canonical_product_name": info["canonical_name"],
        "provisional_sku": provisional_sku(row["product"]) if row["product"] else "",
        "sku_status": "provisional_reference_key_no_vendor_sku_in_source" if row["product"] else "missing_product",
        "product_exists_in_reference_workbook": bool(row["product"]),
        "allocation_method": "",
        "allocation_confidence": "",
        "reference_id": "",
        "reference_year": None,
        "reference_sheet": "",
        "reference_dr_number": "",
        "reference_workbook_rows": "",
        "source_cost_missing": row["total_tp"] is None,
        "parent_financial_anomaly": False,
        "product_assignment_verified": None,
        "record_status": "original_2025_product_row",
        "notes": "Preserved source product row from Sales Report.xlsx.",
    }


# Actual medical lines observed in the supplied workbook. These rows are the only products eligible for allocation.
reference_lines = []
for year, rows in sorted(rows_by_year.items()):
    for row in rows:
        info = class_info(row["product"])
        if replacement_type(row) or info["group"] not in MEDICAL_GROUPS:
            continue
        qty = row["qty"] or 0.0
        net = row["net_cp"] or 0.0
        cost = row["total_tp"] if row["total_tp"] is not None else 0.0
        if qty <= 0 or net <= 0:
            continue
        reference_lines.append({
            **row,
            "product_group": info["group"],
            "product_category": info["category"],
            "classification_confidence": info["confidence"],
            "classification_rule": info["rule"],
            "canonical_product_name": info["canonical_name"],
            "provisional_sku": provisional_sku(row["product"]),
            "net": net,
            "cost": cost,
        })

reference_line_by_delivery = defaultdict(list)
for line in reference_lines:
    reference_line_by_delivery[(line["year"], line["area"].upper(), line["dr_number"].upper(), line["date_delivered"])].append(line)

# Build actual medical transaction baskets using only lines with usable acquisition costs.
groups = defaultdict(list)
for line in reference_lines:
    if line["cost"] <= 0:
        continue
    groups[(line["year"], line["area"].upper(), line["dr_number"], line["date_delivered"])].append(line)

reference_baskets = []
for (year, area_key, dr_number, delivered), lines in groups.items():
    product_lines = {}
    for line in lines:
        product_key = normalized_product(line["product"])
        item = product_lines.setdefault(product_key, {
            "product": line["product"],
            "product_group": line["product_group"],
            "product_category": line["product_category"],
            "classification_confidence": line["classification_confidence"],
            "classification_rule": line["classification_rule"],
            "canonical_product_name": line["canonical_product_name"],
            "provisional_sku": line["provisional_sku"],
            "qty": 0.0,
            "net": 0.0,
            "cost": 0.0,
            "source_rows": [],
        })
        item["qty"] += line["qty"] or 0.0
        item["net"] += line["net"]
        item["cost"] += line["cost"]
        item["source_rows"].append(line["source_row"])
    basket_lines = sorted(product_lines.values(), key=lambda item: (-item["net"], item["product"]))
    total_net = sum(item["net"] for item in basket_lines)
    total_cost = sum(item["cost"] for item in basket_lines)
    if total_net <= 0 or total_cost <= 0:
        continue
    reference_baskets.append({
        "year": year,
        "area": area_key,
        "dr_number": dr_number,
        "date_delivered": delivered,
        "lines": basket_lines,
        "total_net": round(total_net, 2),
        "total_cost": round(total_cost, 2),
        "cost_ratio": total_cost / total_net,
        "product_count": len(basket_lines),
    })

reference_baskets.sort(key=lambda item: (item["area"], item["year"], item["dr_number"], item["date_delivered"]))
for index, basket in enumerate(reference_baskets, start=1):
    basket["reference_id"] = f"REF-XLSX-BASKET-{index:04d}"


def max_products(parent_net):
    if parent_net < 25_000:
        return 2
    if parent_net < 50_000:
        return 3
    if parent_net < 100_000:
        return 4
    if parent_net < 250_000:
        return 6
    if parent_net < 500_000:
        return 8
    if parent_net < 1_000_000:
        return 12
    if parent_net < 5_000_000:
        return 20
    return 30


def desired_products(parent_net):
    if parent_net < 25_000:
        return 1
    if parent_net < 50_000:
        return 2
    if parent_net < 100_000:
        return 3
    if parent_net < 250_000:
        return 4
    if parent_net < 500_000:
        return 6
    if parent_net < 1_000_000:
        return 8
    if parent_net < 5_000_000:
        return 12
    return 16


def select_reference_basket(parent):
    parent_net = parent["net_cp"] or 0.0
    parent_cost = parent["total_tp"] or 0.0
    parent_ratio = parent_cost / parent_net if parent_net else 0.0
    area_key = parent["area"].upper()
    same_area = [basket for basket in reference_baskets if basket["area"] == area_key]
    candidates = same_area or reference_baskets
    cap = max_products(parent_net)
    capped = [basket for basket in candidates if basket["product_count"] <= cap]
    if capped:
        candidates = capped
    desired = desired_products(parent_net)

    def score(basket):
        value_score = abs(math.log(max(basket["total_net"], 0.01) / max(parent_net, 0.01)))
        ratio_score = abs(math.log(max(basket["cost_ratio"], 0.001) / max(parent_ratio, 0.001)))
        recency_score = 0.04 * (2025 - basket["year"])
        size_score = 0.025 * abs(basket["product_count"] - desired)
        return value_score + 2.75 * ratio_score + recency_score + size_score

    basket = min(candidates, key=score)
    return basket, score(basket)


def select_reference_line(parent):
    qty = parent["qty"] or 0.0
    net = parent["net_cp"] or 0.0
    parent_unit = (parent["cp"] or (net / qty if qty else 0.0)) or 0.0
    delivery_key = (2025, parent["area"].upper(), parent["dr_number"].upper(), parent["date_delivered"])
    same_delivery = reference_line_by_delivery.get(delivery_key, [])
    same_area = [line for line in reference_lines if line["area"].upper() == parent["area"].upper()]
    candidates = same_delivery or same_area or reference_lines

    def score(line):
        ref_unit = line["net"] / line["qty"] if line["qty"] else line["cp"] or 0.01
        price_score = abs(math.log(max(ref_unit, 0.01) / max(parent_unit, 0.01))) if parent_unit > 0 else 2.0
        qty_score = 0.15 * abs(math.log(max(line["qty"], 0.01) / max(qty, 0.01))) if qty > 0 else 0.5
        recency_score = 0.04 * (2025 - line["year"])
        delivery_bonus = -2.0 if line in same_delivery else 0.0
        return price_score + qty_score + recency_score + delivery_bonus

    line = min(candidates, key=score)
    return line, score(line), bool(same_delivery)


def allocate_amount(total, weights):
    total = round(float(total or 0.0), 2)
    if not weights:
        return []
    weight_sum = sum(max(float(value or 0.0), 0.0) for value in weights)
    if weight_sum <= 0:
        weights = [1.0] * len(weights)
        weight_sum = float(len(weights))
    values = [round(total * max(float(weight or 0.0), 0.0) / weight_sum, 2) for weight in weights]
    values[max(range(len(values)), key=lambda index: weights[index])] = round(
        values[max(range(len(values)), key=lambda index: weights[index])] + total - sum(values), 2
    )
    return values


full_rows = []
parent_rows = []
audit_rows = []
evidence_index = {}
replacement_counts = Counter()


def register_evidence(product_info, parent_type, reference_id, reference_year, reference_dr, source_rows):
    key = normalized_product(product_info["product"])
    evidence = evidence_index.setdefault(key, {
        "product": product_info["product"],
        "canonical_product_name": product_info["canonical_product_name"],
        "provisional_sku": product_info["provisional_sku"],
        "sku_status": "provisional_reference_key_no_vendor_sku_in_source",
        "product_group": product_info["product_group"],
        "product_category": product_info["product_category"],
        "classification_confidence": product_info["classification_confidence"],
        "classification_rule": product_info["classification_rule"],
        "allocation_count": 0,
        "parent_types": set(),
        "reference_ids": set(),
        "reference_years": set(),
        "reference_dr_numbers": set(),
        "reference_workbook_rows": set(),
    })
    evidence["allocation_count"] += 1
    evidence["parent_types"].add(parent_type)
    evidence["reference_ids"].add(reference_id)
    evidence["reference_years"].add(reference_year)
    evidence["reference_dr_numbers"].add(reference_dr)
    evidence["reference_workbook_rows"].update(f"{reference_year}!{row}" for row in source_rows)


for source in source_2025:
    parent_type = replacement_type(source)
    if not parent_type:
        full_rows.append(normalized_source_row(source))
        continue

    replacement_counts[parent_type] += 1
    parent_id = f"SRC-XLSX-2025-R{source['source_row']:04d}"
    parent_net = round(source["net_cp"] or 0.0, 2)
    parent_total_cp = round(source["total_cp"] or 0.0, 2)
    parent_discount = round(source["discount"] or 0.0, 2)
    parent_cost = round(source["total_tp"] or 0.0, 2)
    parent_gross = round(source["net_income"] if source["net_income"] is not None else parent_net - parent_cost, 2)
    financial_anomaly = parent_net < 0 or parent_cost < 0 or parent_cost > parent_net
    source_cost_missing = source["total_tp"] is None

    parent_record = normalized_source_row(source)
    parent_record.update({
        "replacement_type": parent_type,
        "record_type": "source_parent_replaced",
        "record_status": "preserved_in_parent_lineage_only",
        "parent_financial_anomaly": financial_anomaly,
        "product_assignment_verified": False,
        "notes": "Excluded from the Full 2025 Dataset and replaced by estimated medical-product child rows.",
    })

    child_specs = []
    if parent_type == "blank_or_invalid_product" or parent_cost <= 0:
        line, match_score, same_delivery = select_reference_line(source)
        child_specs = [{
            "product": line["product"],
            "canonical_product_name": line["canonical_product_name"],
            "provisional_sku": line["provisional_sku"],
            "product_group": line["product_group"],
            "product_category": line["product_category"],
            "classification_confidence": line["classification_confidence"],
            "classification_rule": line["classification_rule"],
            "reference_qty": line["qty"],
            "reference_net": line["net"],
            "reference_cost": line["cost"],
            "reference_rows": [line["source_row"]],
            "reference_id": f"REF-XLSX-LINE-{line['year']}-{line['source_row']}",
            "reference_year": line["year"],
            "reference_sheet": str(line["year"]),
            "reference_dr_number": line["dr_number"],
            "allocation_method": "same_delivery_or_nearest_actual_medical_line" if same_delivery else "nearest_actual_medical_line_by_area_price_and_quantity",
            "match_score": match_score,
        }]
        qty_values = [source["qty"] or 0.0]
        net_values = [parent_net]
        total_cp_values = [parent_total_cp]
        discount_values = [parent_discount]
        cost_values = [parent_cost]
        confidence = "low"
        reference_label = child_specs[0]["reference_id"]
        reference_year = child_specs[0]["reference_year"]
        reference_dr = child_specs[0]["reference_dr_number"]
        allocation_method = child_specs[0]["allocation_method"]
        match_score_value = match_score
    else:
        basket, match_score = select_reference_basket(source)
        child_specs = [{
            "product": line["product"],
            "canonical_product_name": line["canonical_product_name"],
            "provisional_sku": line["provisional_sku"],
            "product_group": line["product_group"],
            "product_category": line["product_category"],
            "classification_confidence": line["classification_confidence"],
            "classification_rule": line["classification_rule"],
            "reference_qty": line["qty"],
            "reference_net": line["net"],
            "reference_cost": line["cost"],
            "reference_rows": line["source_rows"],
            "reference_id": basket["reference_id"],
            "reference_year": basket["year"],
            "reference_sheet": str(basket["year"]),
            "reference_dr_number": basket["dr_number"],
            "allocation_method": "nearest_same_channel_actual_medical_basket",
            "match_score": match_score,
        } for line in basket["lines"]]
        net_values = allocate_amount(parent_net, [item["reference_net"] for item in child_specs])
        total_cp_values = allocate_amount(parent_total_cp, [item["reference_net"] for item in child_specs])
        discount_values = allocate_amount(parent_discount, [item["reference_net"] for item in child_specs])
        cost_values = allocate_amount(parent_cost, [item["reference_cost"] for item in child_specs])
        scale = parent_net / basket["total_net"] if basket["total_net"] else 0.0
        qty_values = [round(max(item["reference_qty"] * scale, 0.01), 2) for item in child_specs]
        adjustment = (parent_cost / basket["total_cost"]) / scale if scale and basket["total_cost"] else None
        confidence = "low" if financial_anomaly or match_score > 2.5 or not adjustment or adjustment < 0.5 or adjustment > 2.0 else "medium"
        reference_label = basket["reference_id"]
        reference_year = basket["year"]
        reference_dr = basket["dr_number"]
        allocation_method = "nearest_same_channel_actual_medical_basket"
        match_score_value = match_score

    child_rows = []
    for index, spec in enumerate(child_specs, start=1):
        qty = qty_values[index - 1]
        total_cp = total_cp_values[index - 1]
        discount = discount_values[index - 1]
        net = net_values[index - 1]
        cost = cost_values[index - 1]
        gross = round(net - cost, 2)
        child_id = f"EST-MED-{parent_id}-{index:02d}"
        row = {
            "area": source["area"],
            "dr_number": source["dr_number"],
            "date_delivered": source["date_delivered"],
            "product": spec["product"],
            "qty": qty,
            "cp": round(total_cp / qty, 2) if qty else None,
            "total_cp": total_cp,
            "discount": discount,
            "net_cp": net,
            "tp_unit": round(cost / qty, 2) if qty else None,
            "total_tp": cost,
            "net_income": gross,
            "margin_pct": gross / net if net else None,
            "source_year": 2025,
            "source_sheet": "2025",
            "source_workbook_row": source["source_row"],
            "source_record_id": child_id,
            "parent_source_record_id": parent_id,
            "child_sequence": index,
            "original_product_label": source["product"],
            "replacement_type": parent_type,
            "record_type": "estimated_medical_product_child",
            "product_group": spec["product_group"],
            "product_category": spec["product_category"],
            "classification_confidence": spec["classification_confidence"],
            "classification_rule": spec["classification_rule"],
            "canonical_product_name": spec["canonical_product_name"],
            "provisional_sku": spec["provisional_sku"],
            "sku_status": "provisional_reference_key_no_vendor_sku_in_source",
            "product_exists_in_reference_workbook": True,
            "allocation_method": spec["allocation_method"],
            "allocation_confidence": confidence,
            "reference_id": spec["reference_id"],
            "reference_year": spec["reference_year"],
            "reference_sheet": spec["reference_sheet"],
            "reference_dr_number": spec["reference_dr_number"],
            "reference_workbook_rows": "; ".join(f"{spec['reference_sheet']}!{value}" for value in spec["reference_rows"]),
            "source_cost_missing": source_cost_missing,
            "parent_financial_anomaly": financial_anomaly,
            "product_assignment_verified": False,
            "record_status": "estimated_allocation_not_source_verified",
            "notes": "Product exists in Sales Report.xlsx; assignment to this source parent is estimated from historical sales patterns.",
        }
        child_rows.append(row)
        full_rows.append(row)
        register_evidence(spec, parent_type, spec["reference_id"], spec["reference_year"], spec["reference_dr_number"], spec["reference_rows"])

    child_total_cp = round(sum(row["total_cp"] or 0.0 for row in child_rows), 2)
    child_discount = round(sum(row["discount"] or 0.0 for row in child_rows), 2)
    child_net = round(sum(row["net_cp"] or 0.0 for row in child_rows), 2)
    child_cost = round(sum(row["total_tp"] or 0.0 for row in child_rows), 2)
    child_gross = round(sum(row["net_income"] or 0.0 for row in child_rows), 2)
    expected_gross = round(parent_net - parent_cost, 2)
    reconciliation_status = "PASS" if all(abs(value) <= 0.01 for value in (
        child_total_cp - parent_total_cp,
        child_discount - parent_discount,
        child_net - parent_net,
        child_cost - parent_cost,
        child_gross - expected_gross,
    )) else "FAIL"
    parent_record.update({
        "allocation_method": allocation_method,
        "allocation_confidence": confidence,
        "reference_id": reference_label,
        "reference_year": reference_year,
        "reference_dr_number": reference_dr,
        "source_cost_missing": source_cost_missing,
        "parent_financial_anomaly": financial_anomaly,
        "replacement_child_count": len(child_rows),
    })
    parent_rows.append(parent_record)
    audit_rows.append({
        "parent_source_record_id": parent_id,
        "source_workbook_row": source["source_row"],
        "area": source["area"],
        "dr_number": source["dr_number"],
        "date_delivered": source["date_delivered"],
        "original_product_label": source["product"],
        "replacement_type": parent_type,
        "parent_qty": source["qty"],
        "parent_total_cp": parent_total_cp,
        "parent_discount": parent_discount,
        "parent_net_sales": parent_net,
        "parent_acquisition_cost": parent_cost,
        "parent_gross_profit": parent_gross,
        "child_count": len(child_rows),
        "child_total_cp": child_total_cp,
        "child_discount": child_discount,
        "child_net_sales": child_net,
        "child_acquisition_cost": child_cost,
        "child_gross_profit": child_gross,
        "total_cp_delta": round(child_total_cp - parent_total_cp, 2),
        "discount_delta": round(child_discount - parent_discount, 2),
        "net_delta": round(child_net - parent_net, 2),
        "cost_delta": round(child_cost - parent_cost, 2),
        "gross_delta": round(child_gross - expected_gross, 2),
        "source_gross_vs_net_less_cost_delta": round(parent_gross - expected_gross, 2),
        "reconciliation_status": reconciliation_status,
        "allocation_method": allocation_method,
        "allocation_confidence": confidence,
        "reference_id": reference_label,
        "reference_year": reference_year,
        "reference_dr_number": reference_dr,
        "match_score": match_score_value,
        "source_cost_missing": source_cost_missing,
        "parent_financial_anomaly": financial_anomaly,
        "product_assignment_verified": False,
    })


evidence_rows = []
for evidence in evidence_index.values():
    evidence_rows.append({
        **{key: value for key, value in evidence.items() if not isinstance(value, set)},
        "parent_types": "; ".join(sorted(evidence["parent_types"])),
        "reference_ids": "; ".join(sorted(evidence["reference_ids"])),
        "reference_years": "; ".join(str(value) for value in sorted(evidence["reference_years"])),
        "reference_dr_numbers": "; ".join(sorted(evidence["reference_dr_numbers"])),
        "reference_workbook_rows": "; ".join(sorted(evidence["reference_workbook_rows"])),
    })
evidence_rows.sort(key=lambda row: (row["product_group"], row["product"]))


def total(rows, field):
    return round(sum((row.get(field) or 0.0) for row in rows), 2)


financial_fields = ["total_cp", "discount", "net_cp", "total_tp", "net_income"]
source_totals = {field: total(source_2025, field) for field in financial_fields}
output_totals = {field: total(full_rows, field) for field in financial_fields}
financial_deltas = {field: round(output_totals[field] - source_totals[field], 2) for field in financial_fields}
non_product_output_rows = [row for row in full_rows if not row["product"] or class_info(row["product"])["group"] == "non_product"]

summary = {
    "source_workbook": str(SOURCE),
    "source_sheet": "2025",
    "source_transaction_rows": len(source_2025),
    "unchanged_product_rows": len(source_2025) - len(parent_rows),
    "parent_rows_replaced": len(parent_rows),
    "replacement_counts": dict(replacement_counts),
    "estimated_medical_child_rows": sum(row["child_count"] for row in audit_rows),
    "full_output_rows": len(full_rows),
    "unique_allocated_medical_products": len(evidence_rows),
    "reference_medical_lines": len(reference_lines),
    "reference_medical_baskets": len(reference_baskets),
    "source_totals": source_totals,
    "output_totals": output_totals,
    "financial_deltas": financial_deltas,
    "parent_reconciliation_failures": sum(row["reconciliation_status"] != "PASS" for row in audit_rows),
    "low_confidence_allocations": sum(row["allocation_confidence"] == "low" for row in audit_rows),
    "source_cost_missing_parents": sum(row["source_cost_missing"] for row in audit_rows),
    "negative_margin_parents": sum(row["parent_financial_anomaly"] for row in audit_rows),
    "non_product_output_rows": len(non_product_output_rows),
    "blank_output_products": sum(not row["product"] for row in full_rows),
    "source_quantity": total(source_2025, "qty"),
    "output_quantity": total(full_rows, "qty"),
}

assert len(parent_rows) == 330, f"Expected 330 replacement parents, got {len(parent_rows)}"
assert summary["parent_reconciliation_failures"] == 0, summary
assert all(abs(value) <= 0.01 for value in financial_deltas.values()), summary
assert summary["non_product_output_rows"] == 0, summary
assert summary["blank_output_products"] == 0, summary

payload = {
    "summary": summary,
    "full_rows": full_rows,
    "parent_rows": parent_rows,
    "audit_rows": audit_rows,
    "evidence_rows": evidence_rows,
}
OUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(json.dumps({"summary": summary, "staging_path": str(OUT)}, indent=2))
