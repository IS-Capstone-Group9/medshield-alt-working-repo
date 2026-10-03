from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "docs" / "chapters"
OUT_PATH = OUT_DIR / "MedShield_Chapter_4_Results_and_Discussion_Proponent_Voice.docx"
ASSET_DIR = ROOT / "tmp" / "chapter4_defense" / "assets"

EVIDENCE_PATH = (
    ROOT
    / "outputs"
    / "descriptive_validation"
    / "2026-09-26T19-02-42-874Z"
    / "evidence.json"
)
STL_PATH = (
    ROOT
    / "outputs"
    / "descriptive_validation"
    / "2026-09-26T17-52-02-042Z"
    / "stl.json"
)
ARCHITECTURE_PATH = ROOT / "tmp" / "chapter4" / "assets" / "medshield_architecture.png"
FORECAST_SCREENSHOT = ROOT / "tmp" / "progress_report" / "screenshots" / "forecast-modeling.png"
PLANNING_SCREENSHOT = ROOT / "tmp" / "progress_report" / "screenshots" / "prescriptive-planning.png"
OVERVIEW_SCREENSHOT = ROOT / "tmp" / "progress_report" / "screenshots" / "overview.png"
SALES_SCREENSHOT = ROOT / "tmp" / "progress_report" / "screenshots" / "sales-diagnostics.png"

BLACK = "000000"
NAVY = "17324D"
BLUE = "2E6F95"
ORANGE = "D9951E"
PALE_BLUE = "EAF2F7"
PALE_GRAY = "F4F5F6"
GRID = "D9D9D9"
RED = "A94442"


def peso(value: float) -> str:
    return f"₱{value:,.2f}"


def pct(value: float) -> str:
    return f"{value:.2f}%"


def set_run_font(run, name="Times New Roman", size=None, bold=None, italic=None, color=BLACK):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    rpr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=95, start=115, bottom=95, end=115):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=10)


def add_body(doc, text, first_line=True, keep=False):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.5) if first_line else None
    p.paragraph_format.keep_together = keep
    p.add_run(text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_unnumbered_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, size=14, bold=True)
    return p


def add_source_note(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=8.5, italic=True, color="4D4D4D")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    set_run_font(r, size=10, bold=True)
    return p


def set_picture_alt_text(run, title, description):
    for doc_pr in run._element.xpath(".//wp:docPr"):
        doc_pr.set("title", title)
        doc_pr.set("descr", description)


def add_figure(doc, path, width, caption, source, alt_text, page_break_before=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page_break_before
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    set_picture_alt_text(run, caption, alt_text)
    add_caption(doc, caption)
    add_source_note(doc, source)


def add_table(
    doc,
    caption,
    headers,
    rows,
    widths,
    alignments=None,
    font_size=8.6,
    source=None,
    page_break_before=False,
):
    cap = doc.add_paragraph(style="Caption")
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.page_break_before = page_break_before
    cap.paragraph_format.space_after = Pt(4)
    rr = cap.add_run(caption)
    set_run_font(rr, size=10, bold=True)

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(text))
        set_run_font(r, size=font_size, bold=True, color="FFFFFF")
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cell = cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2:
                set_cell_shading(cell, PALE_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = alignments[idx] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(text))
            set_run_font(r, size=font_size)
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
    if source:
        add_source_note(doc, source)
    else:
        after = doc.add_paragraph()
        after.paragraph_format.space_after = Pt(4)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    p.add_run("• ")
    p.add_run(text)
    return p


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(5)

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for name, size, before, after in (
        ("Heading 1", 13, 12, 5),
        ("Heading 2", 12, 9, 4),
        ("Heading 3", 12, 7, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.font.size = Pt(10)
    caption.font.color.rgb = RGBColor(0, 0, 0)

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.space_after = Pt(3)

    add_page_number(section.footer.paragraphs[0])


def load_font(size, bold=False):
    candidates = [
        "arialbd.ttf" if bold else "arial.ttf",
        "calibrib.ttf" if bold else "calibri.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size)
        except OSError:
            continue
    return ImageFont.load_default()


def chart_base(title, width=1800, height=950):
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((width // 2, 35), title, fill="#000000", font=load_font(36, True), anchor="ma")
    return image, draw


def save_image(image, name):
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / name
    image.save(path, dpi=(220, 220))
    return path


def make_charts(evidence, stl):
    years = evidence["snapshot"]["year_summary"]
    image, draw = chart_base("Annual Net Sales and Workbook Gross Profit")
    left, top, right, bottom = 150, 150, 1720, 800
    maximum = max(max(x["revenue"], x["income"]) for x in years) * 1.1
    for tick in range(0, 6):
        value = maximum * tick / 5
        y = bottom - (bottom - top) * tick / 5
        draw.line((left, y, right, y), fill="#D9D9D9", width=2)
        draw.text((left - 18, y), f"{value / 1_000_000:.0f}", fill="#333333", font=load_font(20), anchor="rm")
    group = (right - left) / len(years)
    for idx, item in enumerate(years):
        center = left + group * (idx + 0.5)
        for offset, key, color in ((-22, "revenue", "#2E86B7"), (22, "income", "#D9A441")):
            value = item[key]
            height = (bottom - top) * value / maximum
            draw.rectangle((center + offset - 18, bottom - height, center + offset + 18, bottom), fill=color)
        draw.text((center, bottom + 18), item["year"], fill="#222222", font=load_font(20), anchor="ma")
    draw.line((left, top, left, bottom), fill="#555555", width=2)
    draw.line((left, bottom, right, bottom), fill="#555555", width=2)
    draw.text((left, top - 28), "Millions of pesos", fill="#222222", font=load_font(19), anchor="la")
    draw.rectangle((left, 100, left + 28, 128), fill="#2E86B7")
    draw.text((left + 40, 114), "Net sales", fill="#222222", font=load_font(21), anchor="lm")
    draw.rectangle((left + 220, 100, left + 248, 128), fill="#D9A441")
    draw.text((left + 260, 114), "Workbook gross profit", fill="#222222", font=load_font(21), anchor="lm")
    annual = save_image(image, "annual_sales_profit.png")

    monthly = evidence["snapshot"]["monthly"]
    image, draw = chart_base("Monthly Net Sales With Missing Periods Preserved")
    left, top, right, bottom = 150, 155, 1720, 790
    values = [x["revenue"] for x in monthly if x["revenue"] is not None]
    maximum = max(values) * 1.08
    for tick in range(0, 6):
        value = maximum * tick / 5
        y = bottom - (bottom - top) * tick / 5
        draw.line((left, y, right, y), fill="#E0E0E0", width=2)
        draw.text((left - 18, y), f"{value / 1_000_000:.0f}", fill="#333333", font=load_font(19), anchor="rm")
    points = []
    for idx, item in enumerate(monthly):
        x = left + (right - left) * idx / (len(monthly) - 1)
        if item["missing"]:
            step = (right - left) / (len(monthly) - 1)
            draw.rectangle((x - step / 2, top, x + step / 2, bottom), fill="#F7E0E0")
            points.append(None)
        else:
            y = bottom - (bottom - top) * item["revenue"] / maximum
            points.append((x, y))
    segment = []
    for point in points + [None]:
        if point is None:
            if len(segment) > 1:
                draw.line(segment, fill="#17324D", width=5, joint="curve")
            segment = []
        else:
            segment.append(point)
    for idx, item in enumerate(monthly):
        if item["period"].endswith("-01"):
            x = left + (right - left) * idx / (len(monthly) - 1)
            draw.text((x, bottom + 18), item["period"][:4], fill="#222222", font=load_font(19), anchor="ma")
    draw.line((left, top, left, bottom), fill="#555555", width=2)
    draw.line((left, bottom, right, bottom), fill="#555555", width=2)
    draw.text((left, 105), "Light red bands mark months with no observed transactions", fill="#6A3535", font=load_font(20), anchor="la")
    draw.text((left, top - 28), "Millions of pesos", fill="#222222", font=load_font(19), anchor="la")
    monthly_chart = save_image(image, "monthly_sales_gaps.png")

    components = stl["segments"][0]["components"]
    image, draw = chart_base("STL Decomposition of Eligible Net Sales From 2020 to 2025", height=1250)
    left, right = 210, 1730
    series = [
        ("observed", "Observed net sales", "#17324D"),
        ("trend", "Trend", "#2E86B7"),
        ("seasonal", "Seasonal component", "#D9951E"),
        ("residual", "Residual", "#777777"),
    ]
    panel_height = 235
    for panel, (key, label, color) in enumerate(series):
        top = 130 + panel * 270
        bottom = top + panel_height
        vals = [x[key] / 1_000_000 for x in components]
        lo, hi = min(vals), max(vals)
        if hi == lo:
            hi += 1
        for tick in range(3):
            y = top + panel_height * tick / 2
            draw.line((left, y, right, y), fill="#E3E3E3", width=2)
        points = []
        for idx, value in enumerate(vals):
            x = left + (right - left) * idx / (len(vals) - 1)
            y = bottom - (value - lo) / (hi - lo) * panel_height
            points.append((x, y))
        draw.line(points, fill=color, width=4, joint="curve")
        draw.text((25, (top + bottom) // 2), label + "\nMillions", fill="#222222", font=load_font(18), anchor="lm")
        draw.text((left - 12, top), f"{hi:.1f}", fill="#555555", font=load_font(16), anchor="ra")
        draw.text((left - 12, bottom), f"{lo:.1f}", fill="#555555", font=load_font(16), anchor="rd")
        draw.line((left, top, left, bottom), fill="#777777", width=2)
        draw.line((left, bottom, right, bottom), fill="#777777", width=2)
    for year_index, year in enumerate(range(2020, 2026)):
        idx = year_index * 12
        x = left + (right - left) * idx / (len(components) - 1)
        draw.text((x, 1220), str(year), fill="#222222", font=load_font(19), anchor="ms")
    stl_chart = save_image(image, "stl_components.png")

    def horizontal_chart(items, label_key, value_key, title, color, name):
        image, draw = chart_base(title)
        left, top, right, bottom = 610, 140, 1700, 830
        items = list(reversed(items[:10]))
        maximum = max(x[value_key] for x in items) * 1.05
        row_h = (bottom - top) / len(items)
        for idx, item in enumerate(items):
            cy = top + row_h * (idx + 0.5)
            bar_len = (right - left) * item[value_key] / maximum
            draw.rectangle((left, cy - 21, left + bar_len, cy + 21), fill=color)
            label = item[label_key]
            if len(label) > 40:
                label = label[:37] + "..."
            draw.text((left - 18, cy), label, fill="#222222", font=load_font(20), anchor="rm")
            draw.text((left + bar_len + 12, cy), f"{item[value_key] / 1_000_000:.1f}", fill="#333333", font=load_font(18), anchor="lm")
        draw.line((left, top, left, bottom), fill="#555555", width=2)
        draw.text(((left + right) // 2, 875), "Net sales in millions of pesos", fill="#222222", font=load_font(21), anchor="ma")
        return save_image(image, name)

    products_chart = horizontal_chart(
        evidence["snapshot"]["top_products"],
        "product",
        "revenue",
        "Top Ten Source Product Labels by Net Sales",
        "#315F7D",
        "top_products.png",
    )
    areas_chart = horizontal_chart(
        evidence["snapshot"]["by_area"],
        "area",
        "revenue",
        "Top Ten Derived Area Labels by Net Sales",
        "#5C8A72",
        "top_areas.png",
    )

    return annual, monthly_chart, stl_chart, products_chart, areas_chart


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    evidence = json.loads(EVIDENCE_PATH.read_text(encoding="utf-8"))
    stl = json.loads(STL_PATH.read_text(encoding="utf-8"))
    annual_chart, monthly_chart, stl_chart, products_chart, areas_chart = make_charts(evidence, stl)

    snapshot = evidence["snapshot"]
    totals = snapshot["totals"]
    summary = evidence["summary"]
    years = snapshot["year_summary"]
    products = snapshot["top_products"]
    areas = snapshot["by_area"]

    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("CHAPTER 4")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run("RESULTS AND DISCUSSION")
    set_run_font(r, size=15, bold=True)

    add_body(
        doc,
        "Chapter 4 reports the results of the MedShield Decision Support System and examines how far the available evidence answers the capstone questions. The clearest result is a traceable historical sales layer connected to a dashboard. A candidate population of 37,728 accepted sales facts covers 2017 to 2025. Missing months remain visible, while product, area, and forecast displays identify their evidence limits. Predictive publication, external-signal forecasting, and operational inventory optimization remain incomplete; the chapter treats evaluated features separately from measured business outcomes.",
    )
    add_body(
        doc,
        "The discussion follows the project objectives and the guiding question: How can MedShield use historical sales data, validated external signals, and scenario-based planning to improve demand visibility and reduce expiry-driven inventory losses for the 2026 planning cycle? Observed sales, candidate mappings, forecast evaluations, weather proxies, and planning scenarios carry different levels of evidence. The distinctions matter because a working screen does not establish an approved model, a verified stock position, or a reduction in expired inventory. Published studies inform the choice of methods, while MedShield receipts provide the evidence for MedShield results (Nguyen et al., 2021; World Health Organization, 2024).",
    )
    add_body(
        doc,
        "The current MedShield sales record ends in December 2025. For elapsed 2026 months without uploaded client sales, the dashboard marks recency-weighted same-calendar values as gap estimates. The estimates support provisional planning while 2026 actual records are pending. They remain separate from observed sales, cannot serve as a 2026 validation holdout, and require reassessment after actual records arrive. Future 2026 months remain unavailable rather than appearing as actual sales.",
    )

    add_heading(doc, "4.1 Presentation of Results", 1)
    add_body(
        doc,
        "Section 4.1 presents the implemented system and the evidence behind its results before discussing the patterns. The objective audit, architecture, and data-quality profile establish which claims can be defended. Historical sales refer to observed 2017 to 2025 records. A 2026 value without an uploaded actual remains a separately labeled estimate pending the client record.",
    )
    add_heading(doc, "4.1.1 Alignment With the Research Questions and Objectives", 2)
    add_body(
        doc,
        "The proposal defined five objectives: describe sales and seasonality, prioritize products and areas, forecast demand with external signals, support inventory decisions, and present the results through five dashboard modules. The implementation moved furthest on the first, second, and fifth objectives. Forecasting has reached controlled evaluation but not release approval. The prescriptive layer can solve constrained scenarios, but the study does not contain the stock, expiry, lead-time, supplier, ordering-cost, or holding-cost data needed to measure an actual reduction in wastage.",
    )
    add_table(
        doc,
        "Table 4.1 Objective achievement audit",
        ["Objective", "Evidence expected", "Observed result", "Assessment"],
        [
            ["SO1", "Sales performance, year-over-year patterns, and STL seasonality", "Historical summaries and a 72-month STL decomposition are available. Ten earlier months remain unobserved, and the full geography mapping still needs owner approval.", "Partially achieved"],
            ["SO2", "Product, territory, and account priorities using Pareto ABC and predictive scoring", "Deterministic Pareto and ABC candidate views are implemented. Product labels are not approved SKUs, and no XGBoost classifier is published.", "Partially achieved"],
            ["SO3", "Territory forecasts using sales, DII, RSI, Prophet, and XGBoost", "Sales-only baselines and offline candidates were evaluated. External joins have zero forecast-eligible rows, and no model is approved or activated.", "Not achieved for operational use"],
            ["SO4", "EOQ, ROP, and allocation decisions aimed at expiry wastage of at most 5 percent", "The interface supports reviewed allocation scenarios. Real inventory and expiry outcomes are absent, so the 5 percent target was not tested.", "Scenario capability only"],
            ["SO5", "Five dashboard modules with upload validation and decision support", "The five principal modules exist. Browser-to-warehouse CSV ingestion remains unfinished; Signal 2 alerts and collaborative filtering are not verified outputs.", "Partially achieved"],
        ],
        widths=[0.75, 1.65, 3.0, 1.1],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER],
        font_size=8.0,
        source="Source: Capstone proposal objectives and the MedShield current-state, analytics, and QA records as of 27 September 2026.",
        page_break_before=False,
    )
    add_body(
        doc,
        "A completed interface does not by itself prove a research outcome. The forecast screen can display a baseline without an approved model, and the planning solver can operate on assumed inputs without verified stock. Each following result therefore pairs an implemented function with the limit of its evidence.",
    )

    add_heading(doc, "4.1.2 Implemented System and Runtime Architecture", 2)
    add_body(
        doc,
        "MedShield was implemented as an authenticated web application. The Next.js frontend displays analytical views and sends requests to an Express and TypeScript gateway. The gateway performs validation, authorization, and Databricks query orchestration. Databricks Unity Catalog stores the analytical source of truth through Bronze, Silver, Audit, and Gold candidate schemas. Supabase supplies identity and audit support; it does not provide the dashboard measures. Python services are used for offline preparation, model experiments, and validation, but they are not a browser fallback.",
    )
    add_body(
        doc,
        "The runtime boundary remains strict. If Databricks or a required view is unavailable, the gateway returns an error and the dashboard shows an unavailable state. Local JSON, seeded values, and old spreadsheets do not replace the missing result. The declared runtime source makes displayed analytical values traceable.",
    )
    add_figure(
        doc,
        ARCHITECTURE_PATH,
        6.25,
        "Figure 4.1 MedShield runtime architecture",
        "Source: MedShield Architecture and Database documentation, September 2026.",
        "Flow from sales and external source data through Databricks Bronze, Silver, and Gold to the authenticated API gateway and Next.js dashboard, with Supabase providing identity.",
    )
    add_table(
        doc,
        "Table 4.2 Implemented system layers",
        ["Layer", "Implemented component", "Responsibility", "Boundary"],
        [
            ["Presentation", "Next.js, React, TypeScript, Chart.js", "Filters, charts, tables, exports, readiness labels, and scenario forms", "Does not query Databricks or external providers directly"],
            ["Gateway", "Express and TypeScript", "Authentication checks, parameter validation, query orchestration, and response contracts", "Server-only credentials; explicit unavailable states"],
            ["Analytical store", "Databricks Unity Catalog", "Bronze source evidence, Silver assessed records, Gold candidate facts and aggregates", "Candidate publication pending business approval"],
            ["Identity and audit", "Supabase Auth and audit persistence", "Session identity, roles, and application audit events", "Not an analytical source"],
            ["Offline analytics", "Python jobs and tests", "Reconciliation, STL, model evaluation, and scenario validation", "Not a live dashboard dependency"],
        ],
        widths=[1.0, 1.65, 2.45, 1.45],
        font_size=8.1,
        source="Source: MedShield Architecture, Database, Security, and Project documentation.",
    )

    add_heading(doc, "4.1.3 Data Sources Coverage and Quality", 2)
    add_body(
        doc,
        "The historical sales candidate covers 2017 to 2025 and contains 37,728 accepted fact rows. The proposal emphasized 2021 to 2025 and three Philippine regions, while the warehouse retains the longer history available in the supplied records. The broader period provides context. Regional interpretation remains conditional on approval of the area mapping. Department of Health and weather sources provide historical context but are not approved sales regressors.",
    )
    add_table(
        doc,
        "Table 4.3 Data sources and current analytical status",
        ["Source", "Coverage", "Current use", "Status and limitation"],
        [
            ["MedShield sales", "2017 to 2025 actuals; 2026 actuals pending", "Historical demand, revenue, product, area, and baseline forecast evaluation", "37,728 accepted facts; elapsed 2026 gaps may be estimated and remain separately labeled"],
            ["DOH surveillance", "2018 to 2026 declared; 2026 partial", "Historical disease context and candidate regressors", "Territory join not approved; no individual health inference"],
            ["PAGASA observations", "2017 to 2024", "Official historical weather context", "No official 2025 observations in the project; no approved sales match"],
            ["NASA POWER", "108 monthly rows for nine areas in 2025", "Separately labeled weather proxy", "Provider proxy, not PAGASA; 32 candidate matches do not approve regression"],
            ["Operational inventory", "Not supplied", "Would support EOQ, ROP, safety stock, allocation, and expiry analysis", "Scenario inputs only"],
        ],
        widths=[1.15, 1.35, 2.35, 1.65],
        font_size=8.1,
        source="Source: MedShield Current State, System Documentation, and publication-readiness evidence.",
    )
    add_body(
        doc,
        "The quality review identified a complete 108-month calendar but only 98 months with observed transactions. The ten unavailable months are May, June, July, August, and October 2017; January and February 2018; and October, November, and December 2019. The gaps indicate missing observations, not zero demand. Null periods prevent trend and validation routines from creating artificial declines.",
    )
    add_table(
        doc,
        "Table 4.4 Sales candidate quality profile",
        ["Quality measure", "Observed value", "Interpretation"],
        [
            ["Accepted fact rows", "37,728", "Checksum-bound candidate population"],
            ["Revenue-eligible rows", "36,848", "Rows included in the net-sales measure"],
            ["Margin-eligible rows", "34,105", "Rows with a valid margin denominator"],
            ["Warning-bearing rows", "32,749", "Warnings remain visible and are not equivalent to rejection"],
            ["Calendar coverage", "108 months total; 98 observed; 10 missing", "Missing periods remain unavailable rather than zero-filled"],
            ["Product identity", "4,561 source labels; 0 approved master entries", "Product findings describe labels, not verified canonical SKUs"],
            ["Area identity", "22 derived labels", "Some labels mix geography, customer type, or business line and require approval"],
            ["Returns evidence", "673 rows", "Signed returns remain separate while contributing to net revenue"],
        ],
        widths=[1.7, 1.65, 3.05],
        font_size=8.3,
        source="Source: Version-bound publication and QA evidence for the 37,728-row Databricks candidate.",
    )

    add_heading(doc, "4.2 Data Analysis", 1)
    add_body(
        doc,
        "Section 4.2 interprets sales, seasonality, product and area priorities, forecast evaluations, and planning scenarios against the research objectives. Recorded performance remains separate from model outputs and user assumptions. Estimates for elapsed 2026 months with pending actual records do not enter the observed 2017 to 2025 totals or serve as forecast-accuracy evidence.",
    )
    add_body(doc, "In the QA review's default October 2025 to September 2026 window, observed sales were ₱61,082,570, estimated gaps were ₱88,804,257, and the combined display was ₱149,886,827, rounded to pesos. The combined amount is a mixed observed-and-estimated display value, not 2026 actual revenue or a forecast validation result. A later client upload can replace an estimate only after the sales pipeline and publication controls are completed (MedShield Capstone Group 9, 2026c).")
    add_heading(doc, "4.2.1 Historical Sales Performance and Seasonality", 2)
    add_body(
        doc,
        f"Across the accepted candidate, observed net sales totaled {peso(totals['total_revenue'])}. Workbook gross profit totaled {peso(totals['total_income'])}. The weighted gross-margin rate was {totals['avg_margin']:.2f} percent, calculated only over the {peso(totals['margin_revenue_basis'])} margin-eligible revenue basis. The transaction-level workbook measure is not company net income. Dividing gross profit by all reported net sales would use the wrong denominator for the stated rate.",
    )
    year_rows = []
    for item in years:
        year_rows.append(
            [
                item["year"],
                f"{item['transactions']:,}",
                peso(item["revenue"]),
                peso(item["income"]),
                peso(item["margin_revenue_basis"]),
            ]
        )
    add_table(
        doc,
        "Table 4.5 Annual sales and workbook gross profit",
        ["Year", "Transactions", "Net sales", "Workbook gross profit", "Margin eligible sales"],
        year_rows,
        widths=[0.65, 1.05, 1.55, 1.55, 1.55],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT],
        font_size=8.2,
        source="Source: Local descriptive calculation receipt for Delta version 15, 26 September 2026. Current-state and QA summaries report matching overall totals for version 16; the version-16 calculation receipt was unavailable in the local checkout. Amounts are candidate values in Philippine pesos.",
    )
    add_figure(
        doc,
        annual_chart,
        6.25,
        "Figure 4.2 Annual net sales and workbook gross profit",
        "Source: Generated from the version-bound descriptive calculation receipt.",
        "Grouped bars showing annual net sales and workbook gross profit from 2017 through 2025. Net sales are highest in 2025, while the gross-profit series uses a separate eligible basis.",
    )
    add_body(
        doc,
        "The annual series records observed activity without proving equal completeness across years. The 2025 revenue increase is material in the candidate dataset, but source adjustments leave product attribution under review. The recorded financial total can be reconciled while product-level and causal explanations remain provisional.",
    )
    add_body(
        doc,
        "From 2024 to 2025, recorded net sales increased from ₱50.26 million to ₱185.40 million, or about 268.9 percent, while workbook gross profit increased from ₱19.56 million to ₱52.33 million, or about 167.5 percent. Gross profit rose substantially, although its increase was smaller in proportion to revenue. The main measurable explanation is the available calculation basis: ₱49.96 million, or 99.4 percent, of 2024 net sales had fields required for the gross-margin calculation, compared with only ₱95.08 million, or 51.3 percent, of 2025 net sales. Within the matched basis, the calculated margin rate was approximately 39.2 percent in 2024 and 55.0 percent in 2025. The chart cannot establish higher costs, heavier discounts, or an unfavorable product mix as the cause. Much of 2025 revenue lacks the same margin-eligible basis, while source adjustments and product attribution remain under review. A causal explanation requires finance approval and complete matched cost evidence.",
    )
    add_figure(
        doc,
        monthly_chart,
        6.25,
        "Figure 4.3 Monthly net sales and missing coverage",
        "Source: Generated from the 108-month Databricks descriptive snapshot.",
        "Monthly net-sales line from 2017 through 2025 with light red bands marking ten months without observed transactions.",
    )
    add_body(
        doc,
        "Observed calendar-month averages were highest in May at ₱9.87 million and June at ₱7.55 million. The pattern calls for seasonal review but does not show a disease or weather effect. Infectious diseases can follow seasonal cycles, yet a shared calendar pattern cannot establish a causal driver (Martinez, 2018). MedShield therefore presents month averages separately from its STL decomposition.",
    )
    seasonality_rows = [
        [item["month"], peso(item["avg_revenue"]), str(item["observed_month_count"])]
        for item in snapshot["seasonality"]
    ]
    add_table(
        doc,
        "Table 4.6 Observed calendar month averages",
        ["Month", "Average net sales", "Observed years"],
        seasonality_rows,
        widths=[1.2, 2.4, 1.5],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER],
        font_size=8.5,
        source="Source: Descriptive snapshot. Method is observed calendar-month average, not STL.",
    )
    add_body(
        doc,
        "The STL routine used a 12-month period, a seasonal window of 13, robust fitting, and no imputation. Early gaps and short runs left only January 2020 through December 2025 eligible. The 72-month decomposition reconstructs the observed series to numerical precision. STL separates trend, seasonal, and remaining variation for inspection (Cleveland et al., 1990); the MedShield plot does not establish a disease or weather cause.",
    )
    add_figure(
        doc,
        stl_chart,
        6.25,
        "Figure 4.4 STL decomposition of eligible net sales",
        "Source: Checksum-bound STL evidence generated with statsmodels 0.14.6.",
        "Four aligned line charts for observed net sales, trend, seasonal component, and residual from January 2020 through December 2025.",
    )

    add_heading(doc, "4.2.2 Product and Area Prioritization", 2)
    add_body(
        doc,
        "The second objective asks which products, territories, and accounts warrant attention. The descriptive layer answers through positive-sales contribution and deterministic Pareto or ABC grouping. No predictive XGBoost classifier has been approved. An ABC class describes contribution within selected history; it does not predict future urgency or stock-out risk.",
    )
    top_product_rows = []
    for rank, item in enumerate(products[:10], start=1):
        top_product_rows.append(
            [str(rank), item["product"], peso(item["revenue"]), pct(item["pct_of_total"]), item["abc"]]
        )
    add_table(
        doc,
        "Table 4.7 Top source product labels by net sales",
        ["Rank", "Source product label", "Net sales", "Share", "ABC"],
        top_product_rows,
        widths=[0.45, 2.65, 1.55, 0.75, 0.55],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
        font_size=8.0,
        source="Source: Delta version-15 descriptive receipt and Databricks product candidate logic. Share denominator is positive eligible product sales of approximately ₱628.57 million, not overall net sales of ₱626.95 million. Labels are not approved canonical SKUs.",
    )
    add_figure(
        doc,
        products_chart,
        6.25,
        "Figure 4.5 Top ten source product labels by net sales",
        "Source: Generated from the descriptive calculation receipt.",
        "Horizontal bars comparing net sales for the ten highest source product labels in the full candidate population.",
    )
    add_body(
        doc,
        f"SPEEDA 2.5IU/0.5ML ranked first with {peso(products[0]['revenue'])}, or {products[0]['pct_of_total']:.2f} percent of positive eligible product sales. The top five labels represented 12.78 percent and the top ten 19.85 percent on the same positive-sales basis. Overall net sales are lower because signed returns contribute to the net figure. The ranked labels show a dispersed portfolio and require a full-population Pareto calculation. Separate SPEEDA entries also suggest possible alias fragmentation. Multiple-criteria ABC research offers a broader prioritization approach, but MedShield currently reports a deterministic revenue-based candidate rather than a validated multi-criteria classifier (Flores and Whybark, 1986).",
    )
    area_rows = []
    for rank, item in enumerate(areas[:10], start=1):
        area_rows.append(
            [
                str(rank),
                item["area"],
                peso(item["revenue"]),
                pct(item["revenue"] / totals["total_revenue"] * 100),
                peso(item["income"]),
            ]
        )
    add_table(
        doc,
        "Table 4.8 Top derived area labels by net sales",
        ["Rank", "Derived label", "Net sales", "Share", "Workbook gross profit"],
        area_rows,
        widths=[0.45, 1.45, 1.55, 0.75, 1.65],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT],
        font_size=8.0,
        source="Source: Databricks area candidate view. Mixed commercial labels; not a pure geography dimension.",
    )
    add_figure(
        doc,
        areas_chart,
        6.25,
        "Figure 4.6 Top ten derived area labels by net sales",
        "Source: Generated from the descriptive calculation receipt.",
        "Horizontal bars comparing net sales for the ten highest derived area labels, led by Government and Hospital.",
    )
    add_body(
        doc,
        "GOVERNMENT contributed 49.55 percent of candidate net sales and HOSPITAL contributed 26.34 percent. The labels describe commercial segments, not geographic territories. Batangas, Quezon, Laguna, Marinduque, Camarines Norte, and Cavite also appear in the legacy area field. The dashboard exposes buyer-sector and geographic mappings separately, but final owner approval remains open. Government-labeled transactions are financially important in the current candidate; the label does not represent a verified territorial market share.",
    )

    add_heading(doc, "4.2.3 Forecast Evaluation and External Signals", 2)
    add_body(
        doc,
        "The third objective was only partly realized. The system can evaluate simple sales-only baselines and persist candidate experiment evidence, but it has no approved operational forecast. In the fixed-origin Government net-sales comparison retained by the dashboard evidence, the last-observed-value method outperformed the seasonal-naive benchmark across the same nine observed holdout months. Both errors were high, and both methods underpredicted on average. The result supports transparent baseline comparison, not model promotion.",
    )
    add_body(
        doc,
        "No 2026 MedShield sales actuals are present in the current Gold contract. The dashboard's 2026 gap estimates cannot serve as a holdout result or be compared with actual 2026 performance. The estimates maintain planning continuity while actual data are pending. Forecast validation remains retrospective and uses observed months within the 2017 to 2025 history.",
    )
    add_table(
        doc,
        "Table 4.9 Fixed origin Government net sales benchmark",
        ["Method", "Paired months", "MAE", "RMSE", "WAPE", "Bias"],
        [
            ["Seasonal naive", "9", "₱2,695,359.26", "₱3,459,768.58", "85.05%", "-₱2,574,060.06"],
            ["Last observed value", "9", "₱1,987,359.74", "₱2,632,312.49", "62.71%", "-₱1,230,712.16"],
        ],
        widths=[1.45, 0.8, 1.25, 1.25, 0.75, 1.2],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT],
        font_size=8.0,
        source="Source: Section 8 historical acceptance evidence. Missing actuals were excluded rather than set to zero.",
    )
    add_body(
        doc,
        "A later offline experiment evaluated 15 configurations: four simple benchmarks, two ETS models, and nine Prophet configurations. Last value remained the development-selected candidate, with 83.20 percent development WAPE and 37.76 percent retrospective 2025 WAPE. No challenger met the promotion criteria, empirical uncertainty intervals were unsupported, and the model release pointer remained empty. Forecasting guidance requires accuracy checks on observations withheld from model fitting (Hyndman and Athanasopoulos, 2021). MedShield's retrospective errors therefore justify withholding an operational forecast even though a candidate method was selected.",
    )
    add_table(
        doc,
        "Table 4.10 Predictive readiness audit",
        ["Evidence item", "Observed result", "Interpretation"],
        [
            ["Candidate configurations", "15 evaluated", "Experiment completed and persisted as candidate evidence"],
            ["Development selection", "Last value; WAPE 83.20%", "Selected among tested candidates but high error"],
            ["Retrospective 2025 audit", "Last value; WAPE 37.76%", "Historical audit only; not a 2026 holdout"],
            ["Intervals", "Unsupported", "Insufficient lead-specific error history for defensible bands"],
            ["DOH and PAGASA forecast rows", "0 forecast eligible", "External regression remains blocked"],
            ["Publication", "No active model release", "No operational forecast may be claimed"],
        ],
        widths=[1.7, 1.8, 2.95],
        font_size=8.3,
        source="Source: MedShield Analytics, Predictive Operations, Current State, and QA records dated 27 September 2026.",
    )
    add_body(
        doc,
        "The external data audit reinforces the forecast boundary. The current candidate contains 2,171 DOH and PAGASA rows, but none are join-ready or forecast-eligible. NASA POWER supplies separately labeled proxy observations and 32 candidate sales matches; the matches do not establish an approved geographic join or predictive benefit. No evidence supports a claim of improved MedShield forecasts from rainfall, disease intensity, or typhoon signals. Current source-identity and publication gates supersede older exploratory validation labels.",
    )

    add_heading(doc, "4.2.4 Prescriptive Planning Results", 2)
    add_body(
        doc,
        "The fourth objective is represented in the system as a controlled planning scenario. The interface can shortlist positive-revenue products, collect user-entered demand, usable stock, protected stock, pack size, cost, supplier limits, and budget, and solve a constrained allocation. Tests cover feasibility, integer packs, stock protection, and a service-first objective. The result is useful for demonstrating how a decision can be structured, but it remains conditional on the values entered by the user.",
    )
    add_table(
        doc,
        "Table 4.11 Operational inputs required for prescriptive use",
        ["Input", "Current status", "Why it matters"],
        [
            ["On-hand and reserved stock", "Not published", "Determines actual availability and stock gaps"],
            ["Expiry dates and batch age", "Not published", "Required to measure expiry-driven wastage"],
            ["Lead time and service level", "Not published", "Required for ROP and safety-stock calculations"],
            ["Ordering and holding cost", "Not published", "Required for defensible EOQ"],
            ["Pack size MOQ and supplier cap", "User-entered assumption", "Constrains feasible orders"],
            ["Budget and approved demand", "User-entered assumption", "Determines the allocation objective and ceiling"],
        ],
        widths=[1.65, 1.55, 3.2],
        font_size=8.3,
        source="Source: MedShield business definitions, analytics guidance, and planning interface.",
    )
    add_body(
        doc,
        "No actual wastage rate can be calculated because no expiry or write-off outcome dataset is available. The proposal's target of expiry-driven wastage at or below 5 percent remains a goal for a future operational evaluation. A claim of achievement would confuse a design objective with a measured result.",
    )

    add_heading(doc, "4.3 Dashboard", 1)
    add_body(
        doc,
        "Section 4.3 explains the dashboard as the visible part of the decision-support system. A reviewer can follow a summary back to its period, source records, evidence label, and model or scenario status. The 2026 display follows the same rule: observed values require published transactions, while elapsed months awaiting actual records may show separately labeled estimates based on earlier years.",
    )
    add_heading(doc, "4.3.1 Dashboard Structure and Navigation", 2)
    add_body(
        doc,
        "The MedShield dashboard is the presentation layer of the decision-support system. An authenticated user enters through Overview and can open Sales Diagnostics, Product Prioritization, Area Prioritization, Forecast Modeling, and Prescriptive Planning. Supporting views show transaction-level sales evidence and weather-source validation. A Data Upload page exists, but client CSV ingestion and publication are still planned. Navigation places the historical result, source record, assumption, and limitation within one review path.",
    )
    add_body(
        doc,
        "A shared top bar carries the selected period and comparison context. Single-year and year-to-year controls draw available years from the dataset. Changing the selection refreshes related cards, charts, and tables for the same population. The header shows the loaded period and dataset state, while missing periods remain gaps rather than zero demand. For an elapsed 2026 month without an uploaded actual, the dashboard may calculate a same-calendar estimate from up to three earlier years, with recency weights of 0.60, 0.30, and 0.10 where observations exist. Chart and KPI labels distinguish estimates from actual sales, and an uploaded actual takes precedence. If an analytical source or required view is unavailable, the interface clears the affected result and shows an unavailable state.",
    )

    add_heading(doc, "4.3.2 Dashboard Modules and Decision Flow", 2)
    add_body(
        doc,
        "Each principal module addresses a planning question. Sales Diagnostics shows recorded performance. The prioritization pages direct further review. Forecast Modeling compares baselines, while Prescriptive Planning examines a user-defined scenario. The workflow supports human review rather than automatic ordering. Figures 4.7 to 4.10 illustrate the interface; the September development captures serve as layout evidence and do not prove current numerical outputs.",
    )
    add_table(
        doc,
        "Table 4.12 Functions of the five principal dashboard modules",
        ["Module", "Question answered", "Main dashboard outputs", "Decision use and boundary"],
        [
            ["Sales Diagnostics", "How did sales, margin, quantity, and seasonality change?", "KPI cards, annual and monthly trends, matched-period growth, heatmap, and STL evidence", "Explains historical patterns; candidate definitions and missing periods remain visible"],
            ["Product Prioritization", "Which product labels contribute most historical value?", "Ranked contribution, focus cohort, Pareto or ABC grouping, and product table", "Focuses product review; it is not SKU approval or an automatic purchase recommendation"],
            ["Area Prioritization", "Which mapped areas or buyer groups require attention?", "Buyer-sector, geographic, and legacy-area shares with traceable mappings", "Directs commercial review; derived mappings and mixed source labels remain provisional"],
            ["Forecast Modeling", "How well do transparent forecasting baselines perform?", "Actual-versus-holdout plots, error metrics, future baseline, and release status", "Supports model comparison; there is no approved champion, interval, or external regression"],
            ["Prescriptive Planning", "What allocation is feasible under entered constraints?", "Priority shortlist, assumption grid, feasibility result, and scenario export", "Supports scenario discussion; it does not issue purchase orders or use verified inventory and expiry data"],
        ],
        widths=[1.15, 1.75, 2.15, 1.45],
        font_size=7.7,
        source="Source: MedShield Current State, Dashboard Design Logic, Analytics, and System Documentation.",
    )
    if OVERVIEW_SCREENSHOT.exists():
        add_figure(doc, OVERVIEW_SCREENSHOT, 6.25, "Figure 4.7 Overview dashboard development capture", "Source: MedShield development interface capture dated 15 September 2026. Layout evidence only; current numerical results require a current version-bound receipt.", "MedShield Overview page with summary cards, period controls, and navigation.")
        add_body(doc, "Overview places the selected period, sales summary, and source status at the beginning of the review. The historical source ends in December 2025. A mixed window extending into 2026 must show observed sales and gap estimates separately; a combined display is not wholly historical sales.")
    if SALES_SCREENSHOT.exists():
        add_figure(doc, SALES_SCREENSHOT, 6.25, "Figure 4.8 Sales diagnostics development capture", "Source: MedShield development interface capture dated 15 September 2026. Layout evidence only; figures and labels require current runtime verification.", "Sales Diagnostics page showing period controls and historical sales charts.")
        add_body(doc, "Sales Diagnostics brings year comparisons and monthly patterns into the same view. Missing months remain gaps rather than zero sales. A reviewer can use the view to locate unusual periods and then inspect the underlying transaction record.")
    if FORECAST_SCREENSHOT.exists():
        add_figure(
            doc,
            FORECAST_SCREENSHOT,
            6.25,
            "Figure 4.9 Forecast modeling development capture",
            "Source: MedShield development interface capture dated 15 September 2026. Historical screen layout only. Any shaded forecast band in the capture is superseded by the current unsupported-interval status; no model is approved for release.",
            "Forecast Modeling page of the MedShield dashboard showing filters, historical actuals, a retrospective holdout, a future baseline, error metrics, and an explicit historical-dataset status.",
        )
        add_body(
            doc,
            "Forecast Modeling places the selected buyer, product, measure, baseline, and horizon above a comparison chart. Actual sales and retrospective holdout predictions appear apart from a future baseline. MAE, RMSE, WAPE, and bias help a reviewer judge the comparison. The older capture shows a band, but the current analytical record does not support an empirical uncertainty interval. No band or future value should be cited as an approved operational forecast.",
        )
    if PLANNING_SCREENSHOT.exists():
        add_figure(
            doc,
            PLANNING_SCREENSHOT,
            6.25,
            "Figure 4.10 Prescriptive planning development capture",
            "Source: MedShield development interface capture dated 15 September 2026. The displayed scope has no eligible positive-revenue products and does not show a solved allocation. Inputs remain scenario assumptions, not purchase orders.",
            "Prescriptive Planning page showing buyer, territory, horizon, budget, and fulfillment controls, an operational-assumption table, and scenario confirmation.",
        )
        add_body(
            doc,
            "Prescriptive Planning begins with scope and budget controls and requests demand, usable stock, protected stock, pack size, cost, and supplier limits. A confirmation control identifies entered values as assumptions rather than verified inventory records. The displayed empty product scope cannot establish solver performance; separate feasibility tests support the scenario capability.",
        )
    add_body(
        doc,
        "View Sales Data supplies transaction details and provenance for inspecting a displayed total. Weather API Validation identifies official observations separately from NASA POWER proxy data. The Data Upload screen is present, but the current upload endpoint rejects submissions with a warehouse-ingestion-required response. File validation, durable ingestion status, and client-controlled publication therefore remain future work. The supporting views help users check provenance, while the upload limitation prevents a claim of completed 2026 actual-data turnover.",
    )

    add_heading(doc, "4.3.3 Evidence Labels and User Controls", 2)
    add_body(
        doc,
        "Evidence labels help prevent overinterpretation. Actual identifies observed sales; candidate identifies data or mappings awaiting approval; estimated identifies a calculated gap value; proxy identifies a substitute external observation; baseline identifies a forecasting comparator; blocked identifies a result unavailable under current data gates; and scenario identifies an output based on user assumptions. Loading, empty, infeasible, and unavailable states appear explicitly. The browser does not train a model, call an external weather provider directly, or write a purchase transaction. The boundaries preserve the dashboard's role as decision support for an authorized reviewer.",
    )

    add_heading(doc, "4.3.4 Dashboard Verification", 2)
    add_body(
        doc,
        "The current QA review reports 70 of 70 publication checks passing for the version-16 historical candidate. The report establishes technical reconciliation at the review date, not client approval of product identities, area mappings, financial definitions, or a model release. The local checkout contains the earlier version-15 calculation receipt used for the detailed Chapter 4 tables; the later version-16 calculation receipt cited by QA was unavailable locally. The QA review also reports 21 backend unit tests, eight descriptive browser cases, five area and buyer-scope regressions, and six matched-period cases passing, with four retired legacy cases skipped. The sets overlap and are not combined into one test total.",
    )
    add_table(
        doc,
        "Table 4.13 Validation and acceptance evidence",
        ["Validation area", "Observed result", "Interpretation"],
        [
            ["Publication readiness", "70 pass 0 fail 0 unverified", "Technical candidate checks passed; business approval still open"],
            ["Descriptive population", "37,728 facts 108 calendar months 98 observed months", "Counts and periods reconcile; ten months remain unavailable"],
            ["STL", "72 reconciled months and four components", "Qualifying 2020 to 2025 segment only"],
            ["Backend unit tests", "21 passed", "Covers current backend logic including filter validation"],
            ["Final descriptive browser run", "8 passed", "Live version-bound descriptive behaviors and exports"],
            ["Area and buyer regression", "5 passed", "Scoped classification behavior"],
            ["Matched-period and growth suite", "6 passed 4 legacy skipped", "Skipped tests are retired and are not counted as passes"],
            ["Predictive warehouse candidate", "1,092 rows and 15 evaluations persisted", "Candidate experiment evidence; active model pointer remains empty"],
        ],
        widths=[1.6, 2.0, 2.9],
        font_size=8.1,
        source="Source: MedShield System QA Review, Current State, Database, and Predictive Operations records dated 27 September 2026.",
    )

    add_heading(doc, "4.4 Summary of Findings", 1)
    add_body(
        doc,
        "Section 4.4 answers the proposal's practical questions using the observed 2017 to 2025 record, tested system behavior, and documented limits. The 2026 gap estimates provide interim planning context while actual transactions are pending. They do not change the historical totals or establish a measured 2026 outcome.",
    )
    add_heading(doc, "4.4.1 Answers to the Practical Questions", 2)
    add_body(doc, "For the question of which products performed best, SPEEDA 2.5IU/0.5ML ranks first among source product labels in the current candidate. The ranking remains provisional because the approved product master is empty and aliases may divide a commercial product. Revenue-based ABC grouping directs review but does not establish clinical importance or future demand.")
    add_body(doc, "For next-season sales, MedShield reports retrospective baseline comparisons but cannot provide an approved operational forecast. The selected last-value method retained large errors, external disease and weather joins were not forecast-eligible, and 2026 actual sales were unavailable for a genuine 2026 holdout. The dashboard's weighted 2026 gap values are estimates for elapsed periods, not validated forecasts.")
    add_body(doc, "For excess or insufficient stock, sales alone cannot establish a stock position. On-hand units, reserved stock, lead times, and expiry dates were not supplied. For possible action, the planning module can test a user-entered allocation scenario and report feasibility. No scenario proves a purchase quantity, avoided expiry, or the proposal's five-percent wastage target without approved operational data.")
    add_heading(doc, "4.4.2 Historical Patterns", 2)
    add_body(
        doc,
        "The technically reconciled candidate shows ₱626.95 million in observed net sales across 2017 to 2025. Recorded activity is uneven by year and by month, with ten missing months in the early history and a large 2025 value requiring careful provenance review. May and June have the highest observed calendar-month averages. STL provides a separate 2020 to 2025 trend, seasonal, and residual decomposition without imputing the earlier gaps.",
    )
    add_heading(doc, "4.4.3 Forecast Comparison", 2)
    add_body(
        doc,
        "On the comparable fixed-origin Government holdout, last observed value performed better than seasonal naive, but its WAPE remained 62.71 percent. In the larger offline experiment, last value also remained selected, with 83.20 percent development WAPE and 37.76 percent retrospective WAPE. Large errors, unsupported uncertainty intervals, and unapproved external joins justify keeping the model unpublished.",
    )
    add_heading(doc, "4.4.4 Product and Area Priorities", 2)
    add_body(
        doc,
        "SPEEDA 2.5IU/0.5ML is the highest-revenue source product label in the current candidate, while GOVERNMENT is the highest-revenue derived area label. Both findings are useful for review, but neither should be treated as a final canonical SKU or a geographic market-share statement. The approved product master is empty, and the legacy area field mixes different commercial meanings.",
    )
    add_heading(doc, "4.4.5 Planning Scenarios", 2)
    add_body(
        doc,
        "The system can produce constrained allocation scenarios after a user supplies operational assumptions. A result identifies feasible allocations under the entered values and exposes infeasible constraints. Real EOQ, ROP, safety stock, and expiry reduction require published inventory and procurement data.",
    )
    add_heading(doc, "4.4.6 Decision Support Value", 2)
    add_body(
        doc,
        "MedShield places source coverage, KPI definitions, historical patterns, priority candidates, forecast errors, and planning assumptions in one review path. The arrangement supports demand review, but no before-and-after measure shows a quantified improvement in visibility. No REC-approved participant evaluation appears in the available record. System tests establish selected behaviors, not perceived usefulness or user acceptance.",
    )

    add_heading(doc, "4.4.7 Discussion of the North Star", 2)
    add_body(
        doc,
        "The evidence supports the demand-visibility aim more directly than the expiry-loss aim. MedShield consolidates historical sales, preserves gaps, shows KPI bases, and displays the status of models and mappings. A reviewer can move from a historical priority to an allocation scenario. The workflow makes assumptions easier to inspect, although no before-and-after study measures a visibility gain.",
    )
    add_body(
        doc,
        "The study does not demonstrate reduced expiry-driven losses. Measurement requires a baseline of expiry or write-off losses, comparable post-implementation observations, and operational inventory inputs. The current system can prepare an evaluation by recording scenario assumptions and decisions. Until outcome data exist, MedShield provides information and controls for planning without proving the financial effect of a plan.",
    )
    add_body(
        doc,
        "Simple models and explicit release gates proved useful in the forecast review. Prophet configurations did not earn promotion merely through greater complexity. Last value performed best among the tested candidates, but large errors and weak uncertainty evidence kept the model unpublished. An authorized user therefore retains the final planning decision, with the basis of any scenario available for inspection.",
    )

    add_heading(doc, "4.4.8 Limitations", 2)
    add_bullet(doc, "The Gold sales objects remain candidate publications pending named business and finance approval.")
    add_bullet(doc, "The product master has no approved entries, so product analysis uses source labels and cannot standardize aliases or pack conversions.")
    add_bullet(doc, "The legacy area field combines geographic, customer-type, and business-line labels; final mapping approval is still required.")
    add_bullet(doc, "Ten months have no observed transactions. The study preserves the gaps but cannot independently determine whether missing files or absent business activity caused them.")
    add_bullet(doc, "DOH and PAGASA rows are not approved for territory-level forecasting, while NASA POWER remains a separately labeled proxy.")
    add_bullet(doc, "No predictive model or uncertainty interval is approved for operational publication.")
    add_bullet(doc, "Inventory, expiry, supplier, lead-time, service-level, ordering-cost, holding-cost, and stock-position data are unavailable.")
    add_bullet(doc, "The study contains technical and browser validation but no completed REC-approved participant evaluation or final client acceptance record.")

    add_heading(doc, "4.4.9 Overall Synthesis", 2)
    add_body(
        doc,
        "MedShield has a working historical decision-support foundation. The system reconciles 37,728 candidate sales facts, presents 2017 to 2025 performance, preserves ten missing months, decomposes a qualifying 72-month series through STL, ranks product and area labels, compares common-holdout baselines, and exposes constrained planning scenarios through an authenticated dashboard architecture. The descriptive and presentation objectives have the strongest evidence.",
    )
    add_body(
        doc,
        "Important limits remain. Product identity and area mappings still need owner approval. External signals are not forecast-eligible. The selected forecast remains unpublished because its errors are high and uncertainty support is incomplete. Prescriptive outputs remain scenarios, and the expiry-wastage target has not been measured. MedShield can therefore be defended as a historical DSS for traceable review, not as proof of an operational forecast or reduced expiry losses.",
    )

    add_unnumbered_heading(doc, "References")
    references = [
        "Cleveland, R. B., Cleveland, W. S., McRae, J. E., & Terpenning, I. (1990). STL: A seasonal-trend decomposition procedure based on Loess. Journal of Official Statistics, 6(1), 3–73.",
        "Flores, B. E., & Whybark, D. C. (1986). Multiple criteria ABC analysis. International Journal of Operations & Production Management, 6(3), 38–46. https://doi.org/10.1108/eb054765",
        "Hyndman, R. J., & Athanasopoulos, G. (2021). Forecasting: Principles and practice (3rd ed.). OTexts. https://otexts.com/fpp3/",
        "Martinez, M. E. (2018). The calendar of epidemics: Seasonal cycles of infectious diseases. PLOS Pathogens, 14(11), e1007327. https://doi.org/10.1371/journal.ppat.1007327",
        "MedShield Capstone Group 9. (2026a). Capstone project proposal objectives and conceptual framework. Unpublished manuscript.",
        "MedShield Capstone Group 9. (2026b). Descriptive calculation and STL receipts, Delta version 15. Internal validation evidence, 26 September 2026.",
        "MedShield Capstone Group 9. (2026c). System QA review and current-state summary. Internal verification records, 27 September 2026.",
        "MedShield Capstone Group 9. (2026d). Analytics, architecture, and predictive operations. Internal project documentation, 27 September 2026.",
        "Nguyen, A., Lamouri, S., Pellerin, R., Tamayo, S., & Lekens, B. (2021). Data analytics in pharmaceutical supply chains: State of the art, opportunities, and challenges. International Journal of Production Research, 60(22), 6888–6907. https://doi.org/10.1080/00207543.2021.1950937",
        "World Health Organization. (2024). Digital transformation handbook for health supply chain architecture. https://www.who.int/publications/i/item/9789240101197",
    ]
    for reference in references:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(6)
        p.add_run(reference)

    # Report completed work in past tense while retaining a clear review-date boundary.
    past_tense_paragraphs = {
        2: "Chapter 4 reported the results of the MedShield Decision Support System and examined how far the available evidence answered the capstone questions. The clearest result was a traceable historical sales layer connected to a dashboard. A candidate population of 37,728 accepted sales facts covered 2017 to 2025. Missing months remained visible, while product, area, and forecast displays identified their evidence limits. Predictive publication, external-signal forecasting, and operational inventory optimization remained incomplete; the chapter treated evaluated features separately from measured business outcomes.",
        3: "The discussion followed the project objectives and the guiding question: How could MedShield use historical sales data, validated external signals, and scenario-based planning to improve demand visibility and reduce expiry-driven inventory losses for the 2026 planning cycle? Observed sales, candidate mappings, forecast evaluations, weather proxies, and planning scenarios carried different levels of evidence. The distinctions mattered because a working screen did not establish an approved model, a verified stock position, or a reduction in expired inventory. Published studies informed the choice of methods, while MedShield receipts provided evidence for MedShield results (Nguyen et al., 2021; World Health Organization, 2024).",
        4: "The MedShield sales record available for the review ended in December 2025. For elapsed 2026 months without uploaded client sales, the dashboard marked recency-weighted same-calendar values as gap estimates. The estimates supported provisional planning while 2026 actual records were pending. They remained separate from observed sales, could not serve as a 2026 validation holdout, and required reassessment after actual records became available. Future 2026 months remained unavailable rather than appearing as actual sales.",
        6: "Section 4.1 presented the implemented system and the evidence behind its results before discussing the patterns. The objective audit, architecture, and data-quality profile established which claims could be defended. Historical sales referred to observed 2017 to 2025 records. A 2026 value without an uploaded actual remained a separately labeled estimate pending the client record.",
        8: "The proposal defined five objectives: describe sales and seasonality, prioritize products and areas, forecast demand with external signals, support inventory decisions, and present the results through five dashboard modules. The implementation advanced furthest on the first, second, and fifth objectives. Forecasting reached controlled evaluation but not release approval. The prescriptive layer solved constrained scenarios, but the study lacked the stock, expiry, lead-time, supplier, ordering-cost, and holding-cost data needed to measure an actual reduction in wastage.",
        11: "A completed interface did not by itself prove a research outcome. The forecast screen displayed a baseline without an approved model, and the planning solver operated on assumed inputs without verified stock. Each result therefore paired an implemented function with the limit of its evidence.",
        13: "MedShield was implemented as an authenticated web application. The Next.js frontend displayed analytical views and sent requests to an Express and TypeScript gateway. The gateway performed validation, authorization, and Databricks query orchestration. Databricks Unity Catalog stored the analytical source of truth through Bronze, Silver, Audit, and Gold candidate schemas. Supabase supplied identity and audit support but did not provide dashboard measures. Python services supported offline preparation, model experiments, and validation; they did not act as a browser fallback.",
        14: "The runtime boundary remained strict. If Databricks or a required view was unavailable, the gateway returned an error and the dashboard showed an unavailable state. Local JSON, seeded values, and old spreadsheets did not replace the missing result. The declared runtime source made displayed analytical values traceable.",
        21: "The historical sales candidate covered 2017 to 2025 and contained 37,728 accepted fact rows. The proposal emphasized 2021 to 2025 and three Philippine regions, while the warehouse retained the longer history available in the supplied records. The broader period provided context. Regional interpretation remained conditional on approval of the area mapping. Department of Health and weather sources provided historical context but were not approved sales regressors.",
        24: "The quality review identified a complete 108-month calendar but only 98 months with observed transactions. The ten unavailable months were May, June, July, August, and October 2017; January and February 2018; and October, November, and December 2019. The gaps indicated missing observations, not zero demand. Null periods prevented trend and validation routines from creating artificial declines.",
        28: "Section 4.2 interpreted sales, seasonality, product and area priorities, forecast evaluations, and planning scenarios against the research objectives. Recorded performance remained separate from model outputs and user assumptions. Estimates for elapsed 2026 months with pending actual records did not enter the observed 2017 to 2025 totals or serve as forecast-accuracy evidence.",
        29: "In the QA review's default October 2025 to September 2026 window, observed sales were ₱61,082,570, estimated gaps were ₱88,804,257, and the combined display was ₱149,886,827, rounded to pesos. The combined amount was a mixed observed-and-estimated display value, not 2026 actual revenue or a forecast validation result. A later client upload could replace an estimate only after completion of the sales pipeline and publication controls (MedShield Capstone Group 9, 2026c).",
        31: "Across the accepted candidate, observed net sales totaled ₱626,946,741.81. Workbook gross profit totaled ₱288,487,498.72. The weighted gross-margin rate was 60.26 percent, calculated only over the ₱478,720,670.03 margin-eligible revenue basis. The transaction-level workbook measure was not company net income. Dividing gross profit by all reported net sales would have used the wrong denominator for the stated rate.",
        33: "Source: Local descriptive calculation receipt for Delta version 15, 26 September 2026. Current-state and QA summaries reported matching overall totals for version 16; the version-16 calculation receipt was unavailable in the local checkout. Amounts were candidate values in Philippine pesos.",
        37: "The annual series recorded observed activity without proving equal completeness across years. The 2025 revenue increase was material in the candidate dataset, but source adjustments left product attribution under review. The recorded financial total could be reconciled while product-level and causal explanations remained provisional.",
        38: "From 2024 to 2025, recorded net sales increased from ₱50.26 million to ₱185.40 million, or about 268.9 percent, while workbook gross profit increased from ₱19.56 million to ₱52.33 million, or about 167.5 percent. Gross profit rose substantially, although its increase was smaller in proportion to revenue. The main measurable explanation was the available calculation basis: ₱49.96 million, or 99.4 percent, of 2024 net sales had fields required for the gross-margin calculation, compared with only ₱95.08 million, or 51.3 percent, of 2025 net sales. Within the matched basis, the calculated margin rate was approximately 39.2 percent in 2024 and 55.0 percent in 2025. The chart did not establish higher costs, heavier discounts, or an unfavorable product mix as the cause. Much of 2025 revenue lacked the same margin-eligible basis, while source adjustments and product attribution remained under review. A causal explanation required finance approval and complete matched cost evidence.",
        42: "Observed calendar-month averages were highest in May at ₱9.87 million and June at ₱7.55 million. The pattern called for seasonal review but did not show a disease or weather effect. Infectious diseases could follow seasonal cycles, yet a shared calendar pattern could not establish a causal driver (Martinez, 2018). MedShield therefore presented month averages separately from its STL decomposition.",
        44: "Source: Descriptive snapshot. The method was an observed calendar-month average, not STL.",
        45: "The STL routine used a 12-month period, a seasonal window of 13, robust fitting, and no imputation. Early gaps and short runs left only January 2020 through December 2025 eligible. The 72-month decomposition reconstructed the observed series to numerical precision. STL separated trend, seasonal, and remaining variation for inspection (Cleveland et al., 1990); the MedShield plot did not establish a disease or weather cause.",
        50: "The second objective asked which products, territories, and accounts warranted attention. The descriptive layer answered through positive-sales contribution and deterministic Pareto or ABC grouping. No predictive XGBoost classifier had been approved. An ABC class described contribution within selected history; it did not predict future urgency or stock-out risk.",
        52: "Source: Delta version-15 descriptive receipt and Databricks product candidate logic. The share denominator was positive eligible product sales of approximately ₱628.57 million, not overall net sales of ₱626.95 million. Labels were not approved canonical SKUs.",
        56: "SPEEDA 2.5IU/0.5ML ranked first with ₱22,466,448.23, or 3.57 percent of positive eligible product sales. The top five labels represented 12.78 percent and the top ten 19.85 percent on the same positive-sales basis. Overall net sales were lower because signed returns contributed to the net figure. The ranked labels showed a dispersed portfolio and required a full-population Pareto calculation. Separate SPEEDA entries also suggested possible alias fragmentation. Multiple-criteria ABC research offered a broader prioritization approach, but MedShield reported a deterministic revenue-based candidate rather than a validated multi-criteria classifier (Flores and Whybark, 1986).",
    }
    past_tense_paragraphs.update({
        62: "GOVERNMENT contributed 49.55 percent of candidate net sales and HOSPITAL contributed 26.34 percent. The labels described commercial segments, not geographic territories. Batangas, Quezon, Laguna, Marinduque, Camarines Norte, and Cavite also appeared in the legacy area field. The dashboard exposed buyer-sector and geographic mappings separately, but final owner approval remained open. Government-labeled transactions were financially important in the candidate; the label did not represent a verified territorial market share.",
        64: "The third objective was only partly realized. The system evaluated simple sales-only baselines and persisted candidate experiment evidence, but no operational forecast had been approved. In the fixed-origin Government net-sales comparison retained by the dashboard evidence, the last-observed-value method outperformed the seasonal-naive benchmark across the same nine observed holdout months. Both errors were high, and both methods underpredicted on average. The result supported transparent baseline comparison, not model promotion.",
        65: "No 2026 MedShield sales actuals were present in the Gold contract reviewed for the study. The dashboard's 2026 gap estimates could not serve as a holdout result or be compared with actual 2026 performance. The estimates maintained planning continuity while actual data were pending. Forecast validation remained retrospective and used observed months within the 2017 to 2025 history.",
        68: "A later offline experiment evaluated 15 configurations: four simple benchmarks, two ETS models, and nine Prophet configurations. Last value remained the development-selected candidate, with 83.20 percent development WAPE and 37.76 percent retrospective 2025 WAPE. No challenger met the promotion criteria, empirical uncertainty intervals were unsupported, and the model release pointer remained empty. Forecasting guidance required accuracy checks on observations withheld from model fitting (Hyndman and Athanasopoulos, 2021). MedShield's retrospective errors therefore justified withholding an operational forecast even though a candidate method was selected.",
        71: "The external data audit reinforced the forecast boundary. The candidate contained 2,171 DOH and PAGASA rows, but none were join-ready or forecast-eligible. NASA POWER supplied separately labeled proxy observations and 32 candidate sales matches; the matches did not establish an approved geographic join or predictive benefit. No evidence supported a claim of improved MedShield forecasts from rainfall, disease intensity, or typhoon signals. The source-identity and publication gates in force at review superseded older exploratory validation labels.",
        73: "The fourth objective was represented in the system as a controlled planning scenario. The interface shortlisted positive-revenue products, collected user-entered demand, usable stock, protected stock, pack size, cost, supplier limits, and budget, and solved a constrained allocation when eligible products and assumptions were supplied. Tests covered feasibility, integer packs, stock protection, and a service-first objective. The result demonstrated how a decision could be structured, but it remained conditional on user-entered values.",
        76: "No actual wastage rate could be calculated because no expiry or write-off outcome dataset was available. The proposal's target of expiry-driven wastage at or below 5 percent remained a goal for a future operational evaluation. A claim of achievement would have confused a design objective with a measured result.",
        78: "Section 4.3 explained the dashboard as the visible part of the decision-support system. A reviewer could follow a summary back to its period, source records, evidence label, and model or scenario status. The 2026 display followed the same rule: observed values required published transactions, while elapsed months awaiting actual records could show separately labeled estimates based on earlier years.",
        80: "The MedShield dashboard was the presentation layer of the decision-support system. An authenticated user entered through Overview and could open Sales Diagnostics, Product Prioritization, Area Prioritization, Forecast Modeling, and Prescriptive Planning. Supporting views showed transaction-level sales evidence and weather-source validation. A Data Upload page existed, but client CSV ingestion and publication remained planned. Navigation placed the historical result, source record, assumption, and limitation within one review path.",
        81: "A shared top bar carried the selected period and comparison context. Single-year and year-to-year controls drew available years from the dataset. Changing the selection refreshed related cards, charts, and tables for the same population. The header showed the loaded period and dataset state, while missing periods remained gaps rather than zero demand. For an elapsed 2026 month without an uploaded actual, the dashboard could calculate a same-calendar estimate from up to three earlier years, with recency weights of 0.60, 0.30, and 0.10 where observations existed. Chart and KPI labels distinguished estimates from actual sales, and an uploaded actual took precedence. If an analytical source or required view was unavailable, the interface cleared the affected result and showed an unavailable state.",
        83: "Each principal module addressed a planning question. Sales Diagnostics showed recorded performance. The prioritization pages directed further review. Forecast Modeling compared baselines, while Prescriptive Planning examined a user-defined scenario. The workflow supported human review rather than automatic ordering. Figures 4.7 to 4.10 illustrated the interface; the September development captures served as layout evidence and did not prove current numerical outputs.",
        88: "Source: MedShield development interface capture dated 15 September 2026. The capture provided layout evidence only; current numerical results required a version-bound receipt.",
        89: "Overview placed the selected period, sales summary, and source status at the beginning of the review. The historical source ended in December 2025. A mixed window extending into 2026 required separate observed-sales and gap-estimate labels; a combined display was not wholly historical sales.",
        92: "Source: MedShield development interface capture dated 15 September 2026. The capture provided layout evidence only; figures and labels required runtime verification.",
        93: "Sales Diagnostics brought year comparisons and monthly patterns into the same view. Missing months remained gaps rather than zero sales. A reviewer could use the view to locate unusual periods and inspect the underlying transaction record.",
        96: "Source: MedShield development interface capture dated 15 September 2026. The capture showed a historical screen layout only. A shaded forecast band appeared in the older capture, but the review record reported unsupported intervals and no approved model release.",
        97: "Forecast Modeling placed the selected buyer, product, measure, baseline, and horizon above a comparison chart. Actual sales and retrospective holdout predictions appeared apart from a future baseline. MAE, RMSE, WAPE, and bias helped a reviewer judge the comparison. The older capture showed a band, but the analytical record did not support an empirical uncertainty interval. No band or future value was presented as an approved operational forecast.",
        100: "Source: MedShield development interface capture dated 15 September 2026. The displayed scope had no eligible positive-revenue products and did not show a solved allocation. Inputs represented scenario assumptions, not purchase orders.",
        101: "Prescriptive Planning began with scope and budget controls and requested demand, usable stock, protected stock, pack size, cost, and supplier limits. A confirmation control identified entered values as assumptions rather than verified inventory records. The displayed empty product scope could not establish solver performance; separate feasibility tests supported the scenario capability.",
        102: "View Sales Data supplied transaction details and provenance for inspecting a displayed total. Weather API Validation identified official observations separately from NASA POWER proxy data. The Data Upload screen was present, but the upload endpoint rejected submissions with a warehouse-ingestion-required response. File validation, durable ingestion status, and client-controlled publication therefore remained future work at the review date. The supporting views helped users check provenance, while the upload limitation prevented a claim of completed 2026 actual-data turnover.",
        104: "Evidence labels helped prevent overinterpretation. Actual identified observed sales; candidate identified data or mappings awaiting approval; estimated identified a calculated gap value; proxy identified a substitute external observation; baseline identified a forecasting comparator; blocked identified a result unavailable under the data gates; and scenario identified an output based on user assumptions. Loading, empty, infeasible, and unavailable states appeared explicitly. The browser did not train a model, call an external weather provider directly, or write a purchase transaction. The boundaries preserved the dashboard's role as decision support for an authorized reviewer.",
        106: "The QA review reported 70 of 70 publication checks passing for the version-16 historical candidate. The report established technical reconciliation at the review date, not client approval of product identities, area mappings, financial definitions, or a model release. The local checkout contained the earlier version-15 calculation receipt used for the detailed Chapter 4 tables; the later version-16 calculation receipt cited by QA was unavailable locally. The QA review also reported 21 backend unit tests, eight descriptive browser cases, five area and buyer-scope regressions, and six matched-period cases passing, with four retired legacy cases skipped. The sets overlapped and were not combined into one test total.",
    })
    past_tense_paragraphs.update({
        110: "Section 4.4 answered the proposal's practical questions using the observed 2017 to 2025 record, tested system behavior, and documented limits. The 2026 gap estimates provided interim planning context while actual transactions were pending. They did not change the historical totals or establish a measured 2026 outcome.",
        112: "For the question of which products performed best, SPEEDA 2.5IU/0.5ML ranked first among source product labels in the candidate. The ranking remained provisional because the approved product master was empty and aliases could divide a commercial product. Revenue-based ABC grouping directed review but did not establish clinical importance or future demand.",
        113: "For next-season sales, MedShield reported retrospective baseline comparisons but could not provide an approved operational forecast. The selected last-value method retained large errors, external disease and weather joins were not forecast-eligible, and 2026 actual sales were unavailable for a genuine 2026 holdout. The dashboard's weighted 2026 gap values were estimates for elapsed periods, not validated forecasts.",
        114: "For excess or insufficient stock, sales alone could not establish a stock position. On-hand units, reserved stock, lead times, and expiry dates were not supplied. For possible action, the planning module tested user-entered allocation scenarios and reported feasibility. No scenario proved a purchase quantity, avoided expiry, or the proposal's five-percent wastage target without approved operational data.",
        116: "The technically reconciled candidate showed ₱626.95 million in observed net sales across 2017 to 2025. Recorded activity was uneven by year and by month, with ten missing months in the early history and a large 2025 value requiring careful provenance review. May and June had the highest observed calendar-month averages. STL provided a separate 2020 to 2025 trend, seasonal, and residual decomposition without imputing the earlier gaps.",
        118: "On the comparable fixed-origin Government holdout, last observed value performed better than seasonal naive, but its WAPE remained 62.71 percent. In the larger offline experiment, last value also remained selected, with 83.20 percent development WAPE and 37.76 percent retrospective WAPE. Large errors, unsupported uncertainty intervals, and unapproved external joins justified keeping the model unpublished.",
        120: "SPEEDA 2.5IU/0.5ML was the highest-revenue source product label in the candidate, while GOVERNMENT was the highest-revenue derived area label. Both findings were useful for review, but neither represented a final canonical SKU or a geographic market-share statement. The approved product master was empty, and the legacy area field mixed different commercial meanings.",
        122: "The system produced constrained allocation scenarios after a user supplied operational assumptions. A result identified feasible allocations under the entered values and exposed infeasible constraints. Real EOQ, ROP, safety stock, and expiry reduction required published inventory and procurement data.",
        124: "MedShield placed source coverage, KPI definitions, historical patterns, priority candidates, forecast errors, and planning assumptions in one review path. The arrangement supported demand review, but no before-and-after measure showed a quantified improvement in visibility. No REC-approved participant evaluation appeared in the available record. System tests established selected behaviors, not perceived usefulness or user acceptance.",
        126: "The evidence supported the demand-visibility aim more directly than the expiry-loss aim. MedShield consolidated historical sales, preserved gaps, showed KPI bases, and displayed the status of models and mappings. A reviewer could move from a historical priority to an allocation scenario. The workflow made assumptions easier to inspect, although no before-and-after study measured a visibility gain.",
        127: "The study did not demonstrate reduced expiry-driven losses. Measurement required a baseline of expiry or write-off losses, comparable post-implementation observations, and operational inventory inputs. The system could have supported a later evaluation by recording scenario assumptions and decisions. Without outcome data, MedShield provided information and controls for planning but did not prove the financial effect of a plan.",
        128: "Simple models and explicit release gates proved useful in the forecast review. Prophet configurations did not earn promotion merely through greater complexity. Last value performed best among the tested candidates, but large errors and weak uncertainty evidence kept the model unpublished. An authorized user therefore retained the final planning decision, with the basis of any scenario available for inspection.",
        130: "The Gold sales objects remained candidate publications pending named business and finance approval.",
        131: "The product master had no approved entries, so product analysis used source labels and could not standardize aliases or pack conversions.",
        132: "The legacy area field combined geographic, customer-type, and business-line labels; final mapping approval remained outstanding.",
        133: "Ten months had no observed transactions. The study preserved the gaps but could not independently determine whether missing files or absent business activity caused them.",
        134: "DOH and PAGASA rows were not approved for territory-level forecasting, while NASA POWER remained a separately labeled proxy.",
        135: "No predictive model or uncertainty interval was approved for operational publication.",
        136: "Inventory, expiry, supplier, lead-time, service-level, ordering-cost, holding-cost, and stock-position data were unavailable.",
        137: "The study contained technical and browser validation but no completed REC-approved participant evaluation or final client acceptance record.",
        139: "MedShield had a working historical decision-support foundation. The system reconciled 37,728 candidate sales facts, presented 2017 to 2025 performance, preserved ten missing months, decomposed a qualifying 72-month series through STL, ranked product and area labels, compared common-holdout baselines, and exposed constrained planning scenarios through an authenticated dashboard architecture. The descriptive and presentation objectives had the strongest evidence.",
        140: "Important limits remained. Product identity and area mappings still needed owner approval. External signals were not forecast-eligible. The selected forecast remained unpublished because its errors were high and uncertainty support was incomplete. Prescriptive outputs remained scenarios, and the expiry-wastage target had not been measured. MedShield could therefore be defended as a historical DSS for traceable review, not as proof of an operational forecast or reduced expiry losses.",
    })
    # Present the results in the proponents' formal third-person voice. Keep the
    # observed figures and the distinction between actual, estimated, and planned work.
    past_tense_paragraphs.update({
        2: "The proponents examined how the MedShield Decision Support System answered the capstone questions using the results available at the time of review. The main result was a dashboard linked to 37,728 accepted sales records from 2017 to 2025. It showed missing months rather than treating them as zero sales and gave users ways to inspect products, areas, and forecasting results. The proponents also identified unfinished work. No forecast had been approved for operational use, external signals had not improved a published model, and inventory scenarios had not yet produced measured business outcomes.",
        3: "The discussion followed the proposal's guiding question: How could MedShield use historical sales, validated external signals, and planning scenarios to improve demand visibility and reduce expiry-driven inventory losses for the 2026 planning cycle? The proponents compared each objective with a result they could verify. Sales records, proposed mappings, model tests, weather proxies, and planning scenarios did not carry the same level of certainty. A working dashboard alone did not prove a reliable forecast or fewer expired products. Published research informed the methods, while MedShield records supplied the evidence for the reported results (Nguyen et al., 2021; World Health Organization, 2024).",
        4: "The latest observed MedShield sales month was December 2025. For elapsed months in 2026 without uploaded client records, the dashboard displayed estimates based on the same calendar months in earlier years. The proponents treated the estimates as temporary planning values while actual sales were pending. They did not include the estimates in historical totals or use them to judge forecast accuracy. Future 2026 months remained unavailable until the needed records or forecasts became available.",
        6: "Section 4.1 presented what the proponents completed and the data used to support the results. The section first compared the five objectives with the available outputs, then described the application and the quality of the sales record. Historical results came from observed transactions dated 2017 to 2025. Any displayed 2026 value without an uploaded actual carried an estimate label.",
        8: "The proposal called for analysis of sales patterns, product and area priorities, demand forecasts with external signals, inventory decisions, and a five-module dashboard. The proponents made the most progress on sales analysis, prioritization, and presentation. They tested forecasting methods but did not approve one for release. The planning module solved scenarios with entered assumptions, although stock, expiry, lead-time, supplier, and cost records were unavailable for measuring a real reduction in wastage.",
        11: "The proponents distinguished a completed screen from a completed research objective. The forecast page showed a baseline, but no approved model. The planning page could solve an allocation, but its inputs were not verified stock records. The objective audit therefore reported what each feature did and what its results could not yet establish.",
        12: "4.1.2 Application and Data Architecture",
        13: "MedShield ran as an authenticated web application. Its Next.js dashboard requested analytical results through an Express and TypeScript gateway, which checked requests and queried Databricks. Databricks Unity Catalog held the sales data and candidate analytical views. Supabase handled identity and audit support, not dashboard sales measures. Python tools supported data preparation and model evaluation outside the browser. The arrangement let the proponents trace a dashboard figure back to its analytical source.",
        14: "The dashboard did not substitute a local spreadsheet or sample value when Databricks was unavailable. The gateway returned an error, and the affected view showed an unavailable state. The proponents kept the behavior visible because a plausible-looking number without a live source would have weakened the sales analysis.",
        18: "Table 4.2 Implemented application layers",
        21: "The sales data available for analysis covered 2017 to 2025 and contained 37,728 accepted records. Although the proposal focused on 2021 to 2025 and three Philippine regions, the proponents retained the earlier records to give the trends more context. Regional conclusions remained provisional because the area mapping had not been approved. Department of Health and weather data supplied background, but none had been accepted as a forecasting input for sales.",
        23: "Source: MedShield Current State, application documentation, and publication-readiness evidence.",
        24: "The calendar spanned 108 months, but transactions appeared in only 98. The missing months were May, June, July, August, and October 2017; January and February 2018; and October, November, and December 2019. The proponents kept the months blank because the available records could not show whether sales were zero or files were missing. Filling the gaps with zero would have distorted the trend and later model tests.",
        28: "Section 4.2 examined what the recorded sales revealed about annual performance, seasonality, product and area priorities, and the forecasting and planning objectives. The proponents separated observed transactions from model estimates and user-entered assumptions. Pending 2026 actuals did not enter the 2017 to 2025 totals or the forecast-accuracy calculations.",
        29: "The dashboard's default October 2025 to September 2026 window illustrated the need for separate labels. It contained ₱61,082,570 in observed sales and ₱88,804,257 in estimates for missing months, for a combined display of ₱149,886,827 after rounding. The combined figure was not actual 2026 revenue. The proponents could replace an estimate with a client sales record only after ingestion and publication checks were completed (MedShield Capstone Group 9, 2026c).",
        31: "The accepted sales data recorded ₱626,946,741.81 in net sales and ₱288,487,498.72 in workbook gross profit. The reported weighted gross-margin rate was 60.26 percent, calculated from ₱478,720,670.03 in sales with enough information for the margin calculation. The proponents did not treat workbook gross profit as company net income. Dividing gross profit by all net sales would not reproduce the reported rate because the margin-eligible sales base was smaller.",
        37: "Annual sales varied across the record, and the 2025 increase stood out. The proponents reported the increase as a recorded result while recognizing differences in source coverage between years. Adjustments in the underlying records also left the product-level explanation for 2025 open to review, even though the overall financial total reconciled.",
        38: "Net sales rose from ₱50.26 million in 2024 to ₱185.40 million in 2025, an increase of about 268.9 percent. Workbook gross profit also rose, from ₱19.56 million to ₱52.33 million, but its proportional increase was smaller at about 167.5 percent. The difference did not, by itself, prove an increase in costs or discounts. In 2024, ₱49.96 million, or 99.4 percent of sales, had the fields needed for a gross-margin calculation. In 2025, only ₱95.08 million, or 51.3 percent, had the same fields. On the matched basis, the calculated margin rate was about 39.2 percent in 2024 and 55.0 percent in 2025. The proponents therefore attributed the gap between the revenue and gross-profit growth rates primarily to unequal calculation coverage in the available records, not to an established change in product mix or profitability. A fuller explanation required matched cost data and finance approval.",
        42: "May had the highest observed calendar-month average at ₱9.87 million, followed by June at ₱7.55 million. The result gave the proponents a reason to examine seasonality, but it did not prove a link with disease outbreaks or weather. Infectious diseases could follow seasonal cycles, yet matching calendar patterns alone did not establish cause (Martinez, 2018). The proponents therefore kept simple month averages separate from the STL analysis.",
        45: "The proponents used STL to examine the part of the sales record with enough uninterrupted observations. The routine used a 12-month period, a seasonal window of 13, robust fitting, and no imputation. Earlier gaps left January 2020 to December 2025 as the eligible 72-month series. The resulting plot separated trend, seasonality, and remaining variation (Cleveland et al., 1990). It helped describe the sales pattern but did not identify disease or weather as its cause.",
        50: "For the product-priority objective, the proponents ranked labels by their contribution to positive sales and grouped them through a deterministic Pareto or ABC method. The rankings showed historical value, not future demand or stock-out risk. No predictive XGBoost classifier had been approved, so the chapter did not present an ABC label as a prediction.",
        56: "SPEEDA 2.5IU/0.5ML led the source product labels with ₱22,466,448.23, or 3.57 percent of positive eligible product sales. The five leading labels contributed 12.78 percent, and the top ten contributed 19.85 percent on the same basis. Signed returns explained why the overall net-sales total was lower than the positive-sales base. Sales were spread across many labels, so the proponents used the full list for Pareto grouping. Separate SPEEDA entries also raised the possibility of naming differences for the same commercial product. Although research described broader multiple-criteria ABC methods, MedShield's reported grouping used revenue and had not been validated as a multi-criteria classifier (Flores and Whybark, 1986).",
        62: "GOVERNMENT accounted for 49.55 percent of candidate net sales, followed by HOSPITAL at 26.34 percent. The proponents did not interpret either label as a province or territory. The legacy area field also contained Batangas, Quezon, Laguna, Marinduque, Camarines Norte, and Cavite, so it mixed buyer groups with places. MedShield showed buyer-sector and geographic mappings separately, but owner approval was still pending. The GOVERNMENT share therefore described the candidate sales grouping, not a verified geographic market share.",
        64: "The proponents completed a forecast comparison but not an operational forecast. In a fixed-origin test using Government net sales, the last-observed-value method performed better than the seasonal-naive method over the same nine observed holdout months. Both methods still had large errors and tended to underpredict. The result helped the proponents compare transparent baselines; it did not justify releasing a forecast for planning.",
        65: "The reviewed Gold data contained no actual MedShield sales for 2026. As a result, the dashboard's 2026 gap estimates could not be used as actual outcomes or a 2026 forecast holdout. The proponents evaluated forecasting methods against observed months from 2017 to 2025 and kept the elapsed-2026 estimates as temporary planning context while client records were pending.",
        68: "The proponents later compared 15 offline configurations: four simple benchmarks, two ETS models, and nine Prophet configurations. Last value had the best development selection result, with 83.20 percent development WAPE and 37.76 percent retrospective 2025 WAPE. No alternative met the release criteria, and the available results did not support empirical uncertainty intervals. The model release pointer consequently remained empty. Forecasting guidance called for evaluation on observations withheld from fitting (Hyndman and Athanasopoulos, 2021). The recorded errors gave the proponents a concrete reason to withhold an operational forecast.",
        71: "External data did not resolve the forecast problem. The reviewed data contained 2,171 DOH and PAGASA rows, but none qualified for an approved sales join or forecast input. NASA POWER contributed separately labeled proxy observations and 32 possible sales matches. The matches did not verify the geographic link or show a predictive benefit. The proponents therefore made no claim of improved MedShield forecasts from rainfall, disease intensity, or typhoon indicators.",
        73: "For inventory planning, the proponents demonstrated a scenario rather than a live ordering decision. The dashboard gathered entered demand, usable and protected stock, pack size, cost, supplier limits, and budget, then solved a constrained allocation when eligible products and assumptions were present. Feasibility, whole-pack, stock-protection, and service-first tests supported the calculation. The output showed what could be allocated under entered conditions; it did not show what the client should purchase from verified stock data.",
        76: "The proponents could not calculate an actual wastage rate because expiry and write-off outcomes were unavailable. The proposal's goal of keeping expiry-driven wastage at or below 5 percent remained untested. Reporting the target as an achieved result would have overstated the evidence.",
        78: "Section 4.3 showed how the MedShield dashboard brought the analyses together for a business reviewer. A user could move from a summary to its period, source records, evidence label, and forecast or scenario status. Observed 2026 sales required published client transactions. When actual records for an elapsed month were pending, the dashboard could show a clearly labeled estimate from earlier years.",
        80: "The dashboard opened with Overview and offered Sales Diagnostics, Product Prioritization, Area Prioritization, Forecast Modeling, and Prescriptive Planning. Supporting pages displayed sales transactions and weather-source checks. The proponents arranged the pages so a reviewer could inspect a result, its source, and its limitations in the same workflow. A Data Upload page was visible, but client CSV ingestion and publication had not been completed.",
        81: "The top bar kept the selected period visible across the dashboard. Single-year and year-to-year controls used years found in the data and refreshed the related cards, charts, and tables together. Missing months stayed blank instead of becoming zero sales. For an elapsed 2026 month without an actual record, the dashboard calculated an estimate from up to three earlier occurrences of the same calendar month, weighted 0.60, 0.30, and 0.10 when observations existed. Labels separated estimates from actual sales, and an uploaded actual took priority. If a required data view failed, the affected display showed an unavailable state rather than an old value.",
        83: "Each main page answered a different business question. Sales Diagnostics described recorded performance; product and area pages pointed to groups worth reviewing; Forecast Modeling compared methods; and Prescriptive Planning tested an entered allocation scenario. The dashboard did not place orders automatically. Figures 4.7 to 4.10 showed the layout captured during development in September. The captures illustrated navigation and controls but did not verify the numerical results available at the later review date.",
        85: "Source: MedShield Current State, Dashboard Design Logic, Analytics, and application documentation.",
        89: "Overview gave the reviewer a starting point: the selected period, a sales summary, and the status of the data source. The observed sales record ended in December 2025. When a selected window extended into 2026, the page needed separate labels for observed sales and estimated gaps. The combined amount could not be read as wholly actual revenue.",
        93: "Sales Diagnostics placed annual comparisons beside monthly patterns. The proponents kept missing months as gaps so a reviewer would not mistake absent transactions for zero demand. The page also offered a path from an unusual period to the underlying sales records.",
        97: "Forecast Modeling showed the selected buyer, product, measure, method, and horizon above its comparison chart. Observed sales, retrospective predictions, and a future baseline occupied distinct parts of the view. MAE, RMSE, WAPE, and bias described how the tested method performed. An older development capture included a shaded band, but the completed analysis did not support an empirical uncertainty interval. The proponents did not present the band or the future baseline as an approved forecast.",
        101: "Prescriptive Planning asked for a scope, budget, demand, usable stock, protected stock, pack size, cost, and supplier limits. A confirmation step marked the entries as assumptions instead of verified inventory. The captured screen had no eligible positive-revenue products, so it did not show a completed allocation. Separate feasibility tests, rather than the screenshot, supported the scenario calculation.",
        102: "View Sales Data let a reviewer inspect the transactions behind a displayed total. Weather API Validation separated official observations from NASA POWER proxy values. Although the Data Upload page existed, its endpoint rejected submissions pending warehouse ingestion. File validation, durable ingestion status, and client-controlled publication remained unfinished. The proponents therefore did not claim completed upload of 2026 actual sales.",
        104: "The proponents used labels to show how much weight a reviewer could place on a result. Actual meant observed sales; candidate meant a value or mapping awaiting approval; estimated meant a calculated replacement for a missing period; proxy marked substitute external observations; baseline marked a forecast comparison; blocked marked a result withheld by data checks; and scenario marked a calculation based on entered assumptions. The dashboard also showed loading, empty, infeasible, and unavailable states. It did not train models in the browser, fetch weather directly from the browser, or create purchase transactions.",
        106: "The QA review recorded 70 passing publication checks for the version-16 historical sales candidate. The result confirmed technical reconciliation at the review date, not business approval of product identities, area mappings, finance definitions, or a forecasting method. The detailed tables in the chapter came from the locally available version-15 calculation receipt; the version-16 calculation receipt cited by QA was not available in the checkout. Other reported checks included 21 backend unit tests, eight descriptive browser cases, five area and buyer-scope regressions, and six matched-period cases. Four retired legacy cases were skipped. The proponents kept the sets separate because some tests overlapped.",
        108: "Source: MedShield QA review, Current State, Database, and Predictive Operations records dated 27 September 2026.",
        110: "Section 4.4 brought the main findings back to the proposal's practical questions. The proponents used observed sales from 2017 to 2025, documented dashboard behavior, and the limits of the available data. Elapsed-2026 estimates gave temporary planning context while actual transactions were pending; they did not establish a measured 2026 result.",
        112: "SPEEDA 2.5IU/0.5ML ranked first among source product labels. The proponents treated the ranking as provisional because the approved product master was empty and different labels could refer to the same commercial product. Revenue-based ABC grouping identified products for closer review, not clinical importance or future demand.",
        113: "MedShield compared retrospective forecasting baselines but did not produce an approved next-season forecast. The selected last-value method still had large errors, external disease and weather joins were not accepted for forecasting, and no 2026 actual sales were available for a genuine 2026 holdout. The weighted values shown for elapsed 2026 gaps remained estimates, not validated predictions.",
        114: "Sales records alone did not show whether stock was excessive or insufficient. The proponents lacked verified on-hand units, reservations, lead times, and expiry dates. The planning page tested allocation scenarios entered by a user and reported whether the constraints could be met. It did not prove a purchase quantity, prevented expiry, or achievement of the proposal's five-percent wastage goal.",
        116: "Observed net sales totaled ₱626.95 million across 2017 to 2025. Activity differed by year and month, ten early months lacked transactions, and the large 2025 total required careful source review. May and June had the highest observed month averages. The proponents used STL only on the uninterrupted 2020 to 2025 series, leaving the earlier missing months unfilled.",
        118: "On the fixed-origin Government holdout, last observed value performed better than seasonal naive, although its WAPE was still 62.71 percent. Last value also led the larger offline comparison, with 83.20 percent development WAPE and 37.76 percent retrospective WAPE. The proponents did not publish a forecasting model because the errors were substantial, uncertainty intervals lacked support, and external data joins remained unapproved.",
        120: "SPEEDA 2.5IU/0.5ML led source product labels by revenue, while GOVERNMENT led the derived area labels. The proponents regarded both as useful starting points for review, not final SKU and geographic market-share findings. The product master had no approved entries, and the older area field mixed places with customer and business categories.",
        122: "MedShield produced constrained allocation results when a user entered the needed assumptions. The output showed feasible quantities under the entered limits or explained why a combination was infeasible. Verified inventory and procurement data were still needed for EOQ, ROP, safety stock, and an assessment of expiry reduction.",
        124: "The dashboard brought sales coverage, KPI definitions, historical patterns, product and area rankings, forecast errors, and planning assumptions into one review path. The proponents considered the arrangement useful for examining demand, but they had no before-and-after measure of improved visibility. The available record also contained no completed REC-approved participant evaluation. Technical tests established selected dashboard behavior, not user acceptance.",
        126: "The strongest result addressed demand visibility. MedShield combined historical sales, visible data gaps, clear KPI bases, and the approval status of mappings and forecasts. A reviewer could move from a sales pattern to a product or area priority and then examine an allocation scenario. The proponents could describe the improved review path, but they could not claim a measured gain in decision quality without a before-and-after study.",
        127: "The expiry-loss aim required a separate measurement. No baseline of expired or written-off stock, later comparison period, or verified inventory record was available. MedShield could record assumptions and allocation decisions for a future evaluation, but the proponents could not show a reduction in expiry losses or an improvement in profit.",
        128: "The forecast work also showed why a simple method deserved careful evaluation. More complex Prophet configurations did not meet the release criteria. Last value led the tested methods, yet its errors were still high and its uncertainty estimates lacked support. The proponents kept the model unpublished and left the final planning decision with an authorized user who could inspect the assumptions behind a scenario.",
        139: "The proponents delivered a working historical business analytics solution. MedShield reconciled 37,728 accepted sales records, displayed 2017 to 2025 performance without hiding ten missing months, examined a 72-month series with STL, ranked product and area labels, compared forecasting baselines on observed holdouts, and tested constrained allocation scenarios. Sales analysis and dashboard presentation had the strongest support among the proposal objectives.",
        140: "Several objectives still needed additional data or approval. Product and area labels were not final, external signals had not qualified for forecasting, and the selected model remained unpublished because its error and uncertainty evidence were insufficient. Allocation outputs depended on entered assumptions, while the expiry-wastage target had not been measured. The proponents could defend MedShield as a historical Decision Support System for business review, but not as proof of an approved operational forecast or reduced expiry losses.",
        148: "MedShield Capstone Group 9. (2026c). QA review and current-state summary. Internal verification records, 27 September 2026.",
    })
    for index, revised_text in past_tense_paragraphs.items():
        paragraph = doc.paragraphs[index]
        assert paragraph.runs, f"Paragraph {index} had no editable text run"
        paragraph.runs[-1].text = revised_text

    past_tense_cells = {
        (0, 1, 2): "Historical summaries and a 72-month STL decomposition were available. Ten earlier months remained unobserved, and the full geography mapping still needed owner approval.",
        (0, 2, 2): "Deterministic Pareto and ABC candidate views were implemented. Product labels were not approved SKUs, and no XGBoost classifier was published.",
        (0, 3, 2): "Sales-only baselines and offline candidates were evaluated. External joins had zero forecast-eligible rows, and no model was approved or activated.",
        (0, 4, 2): "The interface supported reviewed allocation scenarios. Real inventory and expiry outcomes were absent, so the 5 percent target was not tested.",
        (0, 5, 2): "The five principal modules existed. Browser-to-warehouse CSV ingestion remained unfinished; Signal 2 alerts and collaborative filtering were not verified outputs.",
        (0, 3, 3): "Not achieved for operational use at review",
        (1, 1, 3): "Did not query Databricks or external providers directly",
        (1, 3, 3): "Candidate publication awaited business approval",
        (1, 4, 3): "Was not an analytical source",
        (1, 5, 3): "Was not a live dashboard dependency",
        (2, 0, 2): "Use at review",
        (2, 0, 3): "Status and limitation at review",
        (2, 1, 1): "2017 to 2025 actuals; 2026 actuals were pending",
        (2, 1, 3): "37,728 accepted facts; elapsed 2026 gaps could be estimated and remained separately labeled",
        (2, 4, 3): "Provider proxy, not PAGASA; 32 candidate matches did not approve regression",
        (2, 5, 2): "Would have supported EOQ, ROP, safety stock, allocation, and expiry analysis",
        (3, 2, 2): "Rows were included in the net-sales measure",
        (3, 4, 2): "Warnings remained visible and were not equivalent to rejection",
        (3, 5, 2): "Missing periods remained unavailable rather than zero-filled",
        (3, 6, 2): "Product findings described labels, not verified canonical SKUs",
        (3, 7, 2): "Some labels mixed geography, customer type, or business line and required approval",
        (3, 8, 2): "Signed returns remained separate while contributing to net revenue",
        (9, 1, 2): "The experiment was completed and persisted as candidate evidence",
        (9, 2, 2): "The method was selected among tested candidates but retained high error",
        (9, 5, 2): "External regression remained blocked",
        (9, 6, 2): "No operational forecast could be claimed",
        (10, 0, 1): "Status at review",
        (10, 0, 2): "Reason for the input",
        (10, 1, 2): "Determined actual availability and stock gaps",
        (10, 2, 2): "Was required to measure expiry-driven wastage",
        (10, 3, 2): "Was required for ROP and safety-stock calculations",
        (10, 4, 2): "Was required for defensible EOQ",
        (10, 5, 2): "Constrained feasible orders",
        (10, 6, 2): "Determined the allocation objective and ceiling",
        (11, 1, 1): "How had sales, margin, quantity, and seasonality changed?",
        (11, 1, 3): "Explained historical patterns; candidate definitions and missing periods remained visible",
        (11, 2, 1): "Which product labels had contributed most historical value?",
        (11, 2, 3): "Focused product review; did not approve SKUs or recommend purchases automatically",
        (11, 3, 1): "Which mapped areas or buyer groups had required attention?",
        (11, 3, 3): "Directed commercial review; derived mappings and mixed source labels remained provisional",
        (11, 4, 1): "How well had transparent forecasting baselines performed?",
        (11, 4, 3): "Supported model comparison; no champion, interval, or external regression was approved",
        (11, 5, 1): "Which allocation was feasible under entered constraints?",
        (11, 5, 3): "Supported scenario discussion; did not issue purchase orders or use verified inventory and expiry data",
        (12, 1, 2): "Technical candidate checks passed; business approval remained open",
        (12, 2, 2): "Counts and periods reconciled; ten months remained unavailable",
        (12, 4, 2): "Covered backend logic, including filter validation",
        (12, 7, 2): "Skipped tests were retired and were not counted as passes",
        (12, 8, 2): "Candidate experiment evidence persisted; the active model pointer remained empty",
    }
    for (table_index, row_index, cell_index), revised_text in past_tense_cells.items():
        paragraph = doc.tables[table_index].rows[row_index].cells[cell_index].paragraphs[0]
        assert paragraph.runs, f"Table {table_index + 1} row {row_index} cell {cell_index} had no text run"
        paragraph.runs[0].text = revised_text

    doc.core_properties.title = "Chapter 4 Results and Discussion"
    doc.core_properties.subject = "MedShield Decision Support System results and objective audit"
    doc.core_properties.author = "MedShield Capstone Group 9"
    doc.core_properties.keywords = "MedShield, Chapter 4, results, discussion, decision support system"
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build()
