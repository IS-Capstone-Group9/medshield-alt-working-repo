from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs' / 'progress_report'
ASSETS = ROOT / 'tmp' / 'progress_report'
SHOTS = ASSETS / 'screenshots'
OUT.mkdir(parents=True, exist_ok=True)

reference = Image.open(ROOT / 'tmp' / 'progress_reference' / 'page-01.png').convert('RGBA')
reference.crop((105, 48, 250, 195)).save(ASSETS / 'ust-seal.png')
reference.crop((790, 42, 930, 200)).save(ASSETS / 'cics-seal.png')

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.line_spacing = 1.45
styles['Normal'].paragraph_format.space_after = Pt(7)
styles['Normal'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size in [('Title', 20), ('Heading 1', 13), ('Heading 2', 11.5), ('Heading 3', 10.5)]:
    style = styles[name]
    style.font.name = 'Arial'
    style._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(6)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def add_institutional_header():
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.text = ''
    table = header.add_table(rows=1, cols=3, width=Inches(6.85))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.columns[0].width = Inches(1.0)
    table.columns[1].width = Inches(4.85)
    table.columns[2].width = Inches(1.0)
    left, center, right = table.rows[0].cells
    left.paragraphs[0].add_run().add_picture(str(ASSETS / 'ust-seal.png'), width=Inches(0.72))
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    right.paragraphs[0].add_run().add_picture(str(ASSETS / 'cics-seal.png'), width=Inches(0.72))
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp = center.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.line_spacing = 1.0
    for line in ['UNIVERSITY OF SANTO TOMAS', 'College of Information and Computing Sciences', 'Department of Information Systems', '1st Semester AY 2026-2027']:
        run = cp.add_run(line + '\n')
        run.bold = True
        run.font.name = 'Arial'; run.font.size = Pt(9.5)
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            tag = OxmlElement(f'w:{edge}'); tag.set(qn('w:val'), 'nil'); borders.append(tag)
        tcPr.append(borders)

def body(text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0.45)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p

def heading(text, level=1):
    p = doc.add_paragraph(text, style=f'Heading {level}')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = True
    return p

def add_table(headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = str(value)
        set_cell_shading(cell, '1F4E78')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255,255,255); run.font.bold = True; run.font.size = Pt(font_size)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_margins(cell)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if ridx % 2: set_cell_shading(cells[i], 'EAF2F8')
            for p in cells[i].paragraphs:
                p.paragraph_format.line_spacing = 1.05; p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i in (0, len(row)-1) else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs: run.font.size = Pt(font_size)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths): row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table

def figure(filename, caption, width=6.75):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(SHOTS / filename), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.keep_with_next = True
    r = c.add_run(caption); r.bold = True; r.italic = True; r.font.size = Pt(10)

def page_break(): doc.add_page_break()

add_institutional_header()

# Cover page
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(78)
p.paragraph_format.line_spacing = 1.4
p.add_run('A Prescriptive Analytics Driven Decision\nSupport System for Pharmaceutical Demand\nForecasting and Planning in MedShield\nPharmaceutical Corporation')

for gap, text in [
    (45, 'A Capstone Project Progress Report Presented to the\nDepartment of Information Systems\nCollege of Information and Computing Sciences\nUniversity of Santo Tomas'),
    (26, 'In partial fulfillment\nof the requirements for the degree of\nBachelor of Science in Information Systems'),
    (24, 'Angeles, Ethan Gabriel C.\nAtabelo, Kerr Lawrence T.\nBenedicto, Adrian Lei M.\nSan Miguel, Keith Nicolai R.'),
    (17, 'Joseph Richard G. Catubag, MBA'),
    (22, 'September 2026')]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(gap)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text); r.bold = True; r.font.size = Pt(10.5)

page_break()
heading('Chapter 4', 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
heading('Results and Discussion', 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
body('This progress report presents the updated implementation and verification results of the MedShield Pharma Corporation Decision Support System. The current revision advances the dashboard from a static prototype toward a source-backed analytical application. It introduces calendar-aware views, revised product and sales visualizations, explicit buyer clusters, rolling forecast horizons, and scenario-safe prescriptive planning. The system is implemented and testable, but selected data-quality and operational-approval gates remain open.')

heading('4.1 Presentation of Results', 1)
body('The revised DSS now operates through a Next.js dashboard, a TypeScript API gateway, and Python analytical services. The application supports multi-year data from 2017 onward, dynamic calendar behavior, transaction-based buyer classification, descriptive product and territory analysis, baseline forecast validation, and constrained planning scenarios. The changes are recorded in Git commit 4d00c29 on branch ega_medshieldup3.')
body('The current implementation distinguishes descriptive evidence, forecast estimates, and scenario outputs. Actual observations override estimates whenever current-year actual data becomes available. Missing observations remain unavailable rather than being converted to zero. Forecast and prescriptive outputs remain labeled as draft or scenario evidence and are not treated as approved procurement instructions.')

add_table(
    ['Component', 'Implemented progress', 'Current qualification'],
    [
        ['Calendar coverage', 'All Years spans 2017 through the current calendar year and adds the current year automatically.', '2017-2019 Net CP revenue remains unresolved in the local transaction source.'],
        ['Overview and Sales Diagnostics', 'Single Year and Y/Y views use monthly granularity; All Years uses annual granularity.', 'Current-year estimates stop at the current Philippine month and are replaced by actuals when loaded.'],
        ['Product Prioritization', 'Primary Pareto chart ranks products by net sales and overlays cumulative contribution.', 'Y/Y mode is removed; companion trends are monthly for one year and annual for all years.'],
        ['Buyer clustering', 'Government, Private, Internal, and Unknown rules are shared by area, forecast, and planning modules.', 'Classification reflects defined source-label rules and does not prove market ownership beyond the available labels.'],
        ['Forecasting', 'Current-month windows support 3, 6, and 12 months with actual history, backtests, and error bands.', 'Baseline models remain draft; stale sales history is disclosed.'],
        ['Prescriptive planning', 'Source-ranked Pareto shortlist and integer-pack allocation scenario are available.', 'Operational stock, supplier, budget, and policy inputs still require approval.'],
    ], [1.25, 3.15, 2.35], 8.2)

page_break()
heading('4.2 Data Analysis', 1)
body('The updated analytical design follows a consistent time-grain rule. A selected year is examined by month because month-level movement is more useful for identifying seasonality, changes in sales activity, and the timing of planning decisions. All Years is summarized annually to preserve long-term comparability without crowding the interface. Y/Y comparison is retained only where direct month-aligned comparison is analytically meaningful.')
body('Product prioritization now uses the Pareto principle directly. Products are ranked from highest to lowest net sales revenue, shown as bars, while the cumulative contribution is plotted as a percentage line. ABC classes, the product table, the portfolio view, and the companion time trend are recalculated from the same selected period. This prevents the product page from combining an annual ranking with unrelated all-history values.')
body('Buyer classification is now expressed through four mutually exclusive clusters. Government includes national government, public hospitals, and accounts explicitly named as LGUs or government institutions. Private includes generic provincial sales areas not explicitly named as LGUs, hospital and pharmacy accounts, and individual sales accounts. Internal isolates MedShield business labels. Records that do not satisfy these rules remain Unknown.')

add_table(
    ['Buyer cluster', 'Current observed records', 'Interpretation'],
    [
        ['Government', '7,544', 'Explicit national government, public hospital, or LGU-related labels.'],
        ['Private', '22,523', 'Generic provincial, hospital, pharmacy, or individual-account sales labels.'],
        ['Internal', '4,901', 'MedShield administration and internal business-line records.'],
        ['Unknown', '577', 'Records not classifiable under the three defined clusters.'],
    ], [1.35, 1.35, 4.05], 8.5)
body('These counts describe the accepted observed records available to the cluster service during the September 15, 2026 verification. They are classification coverage counts rather than market-share estimates. Revenue and quantity shares are calculated only within the selected cluster and analytical scope.')

heading('4.2.1 Financial Data Limitation', 2)
body('The local records for 2017, 2018, and 2019 currently contain zero values in the approved Net CP field used for net sales revenue. Gross sales fields exist, but substituting them would change the approved financial definition. The revised interface therefore does not certify reconstructed net revenue for those years. The historical ingestion and year-specific financial mappings must be repaired and reconciled before the missing values can be presented as validated revenue.')

page_break()
heading('4.3 Dashboard', 1)
body('The dashboard integrates descriptive, predictive, and prescriptive views into a common interface. The revision emphasizes analytical consistency: time filters change the underlying aggregation, buyer clusters use one shared rule, actual observations remain distinct from estimates, and scenario recommendations display their evidence and limitations.')

heading('4.3.1 Executive Overview', 2)
body('The Executive Overview summarizes cumulative revenue, current-year outlook, peak periods, and leading customer context. Its main financial chart changes grain according to the topbar selection. A selected year displays January through December, Y/Y comparison aligns both selected years by calendar month, and All Years displays annual totals from 2017 through the current year.')
figure('overview.png', 'Figure 4.1 Updated MedShield DSS Executive Overview')
body('The 2026 annual value is labeled as an estimate through the current month. The current month is determined in Philippine time and advances automatically as the calendar changes. When actual monthly data is loaded, it replaces the corresponding estimate while later unavailable months remain outside the displayed current-year evidence.')

page_break()
heading('4.3.2 Sales Diagnostics', 2)
body('Sales Diagnostics applies the same time-grain rules to net sales, gross profit, growth, and margin analysis. Single Year displays monthly values, Y/Y compares matching months between the selected years, and All Years retains an annual comparison. This design reduces the visual bulk of repeated annual bars and makes within-year changes easier to interpret.')
figure('sales-diagnostics.png', 'Figure 4.2 Updated Sales Diagnostics with Dynamic Time Grain')
body('Percentage growth is accompanied by nominal peso movement, and missing months are not treated as zero sales. Gross margin is derived from aggregate gross profit divided by aggregate net sales for the selected scope. Labels distinguish revenue, recorded gross profit, rates, and delivered source units to reduce metric ambiguity.')

page_break()
heading('4.3.3 Product Prioritization', 2)
body('Product Prioritization no longer offers Y/Y comparison because the module is intended to rank the product portfolio rather than compare two annual series. Its primary visualization is a Pareto chart: net sales revenue appears as descending bars and cumulative revenue contribution appears as a percentage line on a secondary axis.')
figure('product-prioritization.png', 'Figure 4.3 Product Revenue Pareto and ABC Portfolio Views')
body('For a selected year, the companion trend chart shows monthly revenue for the leading products. Under All Years, the same chart changes to yearly granularity. The ABC distribution, product portfolio, cumulative curve, and ranking table are recalculated from the selected period so the page presents one consistent analytical population.')

page_break()
heading('4.3.4 Area Prioritization', 2)
body('Area Prioritization separates buyer ownership from geography and customer channel. Users can select Government, Private, Internal, or Unknown, then examine revenue share and product-specific quantity share by geography or channel. The selected year and product filters determine the common population used by the chart and table.')
figure('area-prioritization.png', 'Figure 4.4 Buyer Cluster and Geographic Distribution Analysis')
body('The classification rule gives precedence to explicit institutional wording. A province name alone is treated as a private sales area, while an explicitly named provincial government or LGU is Government. MedShield administration and similar business-line labels are Internal. The Unknown group remains visible so unclassified records are not silently reassigned.')

page_break()
heading('4.3.5 Forecast Modeling', 2)
body('Forecast Modeling now begins with the current Philippine calendar month rather than the month immediately following the latest historical record. The horizon selector provides 3, 6, and 12 months, and the current month counts as the first displayed month. On September 15, 2026, the corresponding windows are September-November 2026, September 2026-February 2027, and September 2026-August 2027.')
figure('forecast-modeling.png', 'Figure 4.5 Current Month Forecast and Historical Validation')
body('The chart retains actual observations, retrospective holdout predictions, future baseline forecasts, and empirical historical-error ranges. Missing calendar months remain visible as gaps. The service discloses the training origin and age of the available sales history. With the current source ending in December 2025, the September 2026 forecast is explicitly based on stale closed-month history rather than being presented as a contemporaneous actual.')
body('Seasonal naive and last-observed-value models remain draft benchmarks. The displayed MAE, RMSE, WAPE, bias, holdout coverage, and forecast origin support evaluation but do not establish an approved champion model or operational forecast accuracy.')

page_break()
heading('4.3.6 Prescriptive Planning', 2)
body('Prescriptive Planning uses the same buyer-cluster rules as Area Prioritization and Forecast Modeling. It ranks products by positive net sales over the latest available 12-month period, selects the top 20 percent capped at five products, and prepares an integer-pack allocation scenario.')
figure('prescriptive-planning.png', 'Figure 4.6 Buyer Scoped Prescriptive Planning Scenario')
body('The optimization objective prioritizes fulfilled-demand fractions and then minimizes cost at the best attainable service level. Users must supply demand, usable stock, protected stock, pack size, price, supplier limits, budget, and minimum fulfillment assumptions. Editing an input invalidates the prior output. Results remain scenario-only until inventory, supplier, capacity, expiry, and policy inputs are approved.')

page_break()
heading('4.3.7 Data Upload Page', 2)
body('The Data Upload page retains the source-import controls and ingestion status log. It supports continued loading of sales and structured analytical data without hard-coding the final year. Current-year actuals are designed to supersede estimates when new monthly records arrive.')
figure('data-upload.png', 'Figure 4.7 Updated Data Upload and Ingestion Status Page')
body('The upload workflow remains subject to validation, duplicate handling, financial reconciliation, and source lineage. The unresolved 2017-2019 Net CP issue must be corrected in the ingestion and mapping layer before a refreshed snapshot can be described as financially validated.')

page_break()
heading('4.4 Summary of Findings', 1)
body('The current progress establishes a functioning analytical dashboard with source-aware controls rather than a static design demonstration. The most significant improvements are the consistent monthly and annual time grains, the product Pareto view, current-year estimates that can be replaced by actuals, the four-cluster buyer model, current-month forecast windows, and the alignment of prescriptive planning with the same cluster definitions.')
body('The results also demonstrate the importance of preserving limitations. The 2017-2019 revenue fields remain unresolved, current forecasts are based on sales history ending in December 2025, product identities and units are not yet fully standardized, and prescriptive outputs lack approved operational inputs. These limitations prevent the system from being described as an operationally certified procurement tool, but they do not negate the implementation and verification of the DSS workflow.')

add_table(
    ['Finding', 'Evidence', 'Status'],
    [
        ['Dynamic calendar behavior', '2017-current All Years range; current-year cutoff advances by Philippine month.', 'Implemented'],
        ['Product analytical method', 'Revenue Pareto bars, cumulative percentage line, ABC and scoped trend.', 'Implemented'],
        ['Buyer segmentation', 'Government, Private, Internal, and Unknown shared across three DSS modules.', 'Implemented'],
        ['Forecast horizon', 'Current month plus selectable 3, 6, or 12 month window.', 'Implemented as draft benchmark'],
        ['Prescriptive decision support', 'Pareto shortlist and constrained integer-pack allocation.', 'Scenario only'],
        ['Historical revenue', 'Net CP is zero for 2017-2019 in the current local transaction source.', 'Blocked pending source repair'],
    ], [2.15, 3.45, 1.15], 8.2)

page_break()
heading('4.5 Refinement of Analytical Outputs Based on Technical Consultation', 1)
body('The latest revision implements the principal recommendations documented during the September 2026 technical consultation. Revenue and gross-profit comparisons now use clearer time grains; growth views retain nominal and percentage evidence; product prioritization uses a direct Pareto representation; buyer behavior is separated into explicit analytical clusters; forecast views include actual and retrospective evidence; and prescriptive output is constrained to a focused shortlist.')
body('The product module was refined further after implementation review. Y/Y comparison was removed from Product Prioritization, the primary chart was converted to a true revenue Pareto chart, and the monthly or annual time trend was retained as a companion visualization. This separates portfolio concentration from time-series behavior while keeping both views under the same selected period.')
body('Forecasting was also revised from a static future year to a rolling calendar window. The forecast begins with the current month and offers 3, 6, and 12 month horizons. Closed-month observations remain the training population, current and future source rows are excluded from training, and stale-history warnings remain visible. This supports planning review without disguising the difference between the current calendar and the latest available actual.')
body('Sector segmentation now follows the project-defined interpretation of source labels. Explicit government and LGU wording takes priority; generic province, hospital, pharmacy, and individual-account labels are assigned to Private; MedShield business lines are assigned to Internal; and unmatched records remain Unknown. The rule is implemented once in the analytical service and reused by Area Prioritization, Forecast Modeling, and Prescriptive Planning.')

heading('4.5.1 Verification Record', 2)
add_table(
    ['Verification activity', 'Observed result'],
    [
        ['Backend analytical regression suite', '29 tests passed across buyer clustering, forecast validation, external regression, and prescriptive planning.'],
        ['Frontend type validation', 'TypeScript no-emit validation passed after the dashboard revisions.'],
        ['Targeted browser verification', 'Four cluster and forecast cases passed; the dedicated Pareto and granularity case also passed.'],
        ['Live service check', 'Frontend and backend services ran locally; forecast start returned 2026-09 and buyer-cluster counts reconciled to the current observed population.'],
        ['Version control evidence', 'Commit 4d00c29 was pushed to the ega_medshieldup3 branch on GitHub.'],
    ], [2.4, 4.35], 8.5)

heading('4.5.2 Remaining Work', 2)
body('The next priority is to repair and reconcile year-specific financial extraction for 2017-2019 without substituting gross sales for Net CP. After correction, the dashboard snapshot, annual totals, monthly product views, and downstream forecast populations should be regenerated from one approved source. Product alias and unit mappings, current inventory, supplier constraints, and authenticated acceptance evidence should then be completed before operational claims are made.')

path = OUT / 'MedShield_Chapter_4_Progress_Report_September_2026.docx'
doc.save(path)
print(path)
